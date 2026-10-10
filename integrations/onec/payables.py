"""Supplier debt (кредиторлик): 1C against SAP B1, supplier by supplier — read-only.

The Director's task of 08.10.2026: "1C билан SAP B1даги кредиторлик
қарздорлигини солиштириб чиқинг — қайси контрагентлар ва суммаларда фарқ
бор, сабаби нима". Done under his written temporary permission of 09.10
(``docs/access-review-2026-10-07.md``, section 4); since 2026-10-10 he can
also ask OPS Manager Bot (the analyst's ``supplier_debt_compare`` tool).

Reads, never writes:
  * 1C — over OData with the GET-only ``OneCClient``: balances of 6010 / 6015
    (payables to suppliers, so'm and currency) and 4310 / 4315 (advances paid
    to suppliers), by counterparty and contract, with each counterparty's ИНН;
  * SAP — the suppliers' balances (OCRD, CardType 'S'): pushed by the
    gateway's ``get_supplier_balances`` (2026-10-09), or a file exported from
    SAP B1, as Excel or CSV. Columns are found by their header (SAP's English
    or Russian export names, or the field names — which is what the gateway
    sends).

Gives the comparison (``compare``), a short text of it for the analyst
(``summary``), and one Excel workbook for the Director and accounting
(``write_workbook`` / ``workbook_bytes``): the comparison, what is only in 1C,
what is only in SAP, the 1C detail by contract, and how to read the reasons.
Nothing goes into the database, 1C or SAP, and no figure is logged — only
counts. The command line is ``scripts/ap_reconcile.py``.
"""

from __future__ import annotations

import csv
import io
import json
import re
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import IO, Any

PAYABLE = ("6010", "6015")    # расчеты с поставщиками и подрядчиками (сўм, валюта)
ADVANCE = ("4310", "4315")    # авансы выданные (supplier advances)
TOLERANCE_SOM = 1000.0
UNOFFSET = "1C'да бир вақтда қарз ҳам, аванс ҳам бор — аванс зачет қилинмаган бўлиши мумкин"
TOLERANCE_SHARE = 0.005       # 0.5 %

# SAP export headers -> our field (lower-cased; English UI, Russian UI, field names)
SAP_HEADERS = {
    "code": ("cardcode", "bp code", "код бп", "код партнера", "код делового партнера", "код"),
    "name": ("cardname", "bp name", "имя бп", "наименование бп", "наименование", "название", "имя"),
    "inn": ("lictradnum", "federal tax id", "инн", "tin", "stir", "стир", "налоговый номер"),
    "balance_sys": ("balancesys", "balance (sc)", "сальдо (св)", "баланс (св)", "account balance (sc)"),
    "balance_lc": ("balance", "account balance", "сальдо", "баланс", "сальдо счета"),
    "balance_fc": ("balancefc", "balance (fc)", "сальдо (ин. вал.)", "баланс (ив)"),
    "currency": ("currency", "валюта", "curr"),
    "card_type": ("cardtype", "bp type", "тип бп", "тип"),
}
_LEGAL = re.compile(r"\b(ооо|мчж|mchj|llc|ltd|xk|хк|ип|ак|ao|ао|оао|зао|чп|xususiy|корхонаси|корхона|фирма|firm|"
                    r"компания|company|group|груп|sp|сп|qk|кк|dk|дк|unitar|unitary|davlat)\b")


# ------------------------------------------------------------------ helpers


def digits(value: Any) -> str:
    return re.sub(r"\D", "", str(value or ""))


def norm_name(name: Any) -> str:
    """A counterparty's name for matching: Cyrillic, lower case, no legal form or punctuation."""
    from integrations.common.translit import latin_to_cyrillic

    text = str(name or "").lower().replace("ё", "е")
    text = latin_to_cyrillic(text) if re.search(r"[a-z]", text) else text
    text = re.sub(r"[\"'«»“”‘’`.,()\-–—/\\]+", " ", text)
    text = _LEGAL.sub(" ", text)
    return " ".join(text.split())


def num(value: Any) -> float:
    if value in (None, ""):
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).replace(" ", "").replace(" ", "")
    if text.count(",") == 1 and text.count(".") == 0:
        text = text.replace(",", ".")
    text = text.replace(",", "")
    try:
        return float(text)
    except ValueError:
        return 0.0


# ------------------------------------------------------------------ 1C


@dataclass
class OneCParty:
    key: str
    name: str
    inn: str
    payable: float = 0.0          # 6010 + 6015, so'm (credit balance = we owe)
    payable_fc: dict[str, float] = field(default_factory=dict)  # 6015 in its currency
    advance: float = 0.0          # 4310 + 4315, so'm (debit = we paid ahead)
    lines: list[dict[str, Any]] = field(default_factory=list)

    @property
    def unoffset(self) -> bool:
        """Debt and an advance on the same supplier at once: the advance may never have been offset (зачет)."""
        return self.payable > TOLERANCE_SOM and self.advance > TOLERANCE_SOM

    @property
    def net(self) -> float:
        """What we owe after the advances: positive = we owe the supplier."""
        return self.payable - self.advance


def onec_parties(data: dict[str, Any]) -> dict[str, OneCParty]:
    """1C's register rows → one record per counterparty (only counterparties, not persons)."""
    accounts = {a["Ref_Key"]: (str(a.get("Code", "")).strip(), a.get("Description", "")) for a in data["chart"]}
    parties = {p["Ref_Key"]: p for p in data["parties"]}
    contracts = {c["Ref_Key"]: c.get("Description", "") for c in data.get("contracts", [])}
    currencies = {c["Ref_Key"]: c.get("Description", "") for c in data.get("currencies", [])}
    found: dict[str, OneCParty] = {}
    for row in data["rows"]:
        code, account_name = accounts.get(row.get("Account_Key"), ("", ""))
        if code not in PAYABLE + ADVANCE or not str(row.get("ExtDimension1_Type", "")).endswith("Catalog_Контрагенты"):
            continue
        key = row.get("ExtDimension1") or ""
        info = parties.get(key, {})
        party = found.setdefault(key, OneCParty(key, info.get("Description") or "(номи йўқ)", digits(info.get("ИНН"))))
        cr, dr = num(row.get("СуммаBalanceCr")), num(row.get("СуммаBalanceDr"))
        if code in PAYABLE:
            party.payable += cr - dr
            if code == "6015":
                cur = currencies.get(row.get("Валюта_Key"), "валюта")
                party.payable_fc[cur] = party.payable_fc.get(cur, 0.0) + num(row.get("ВалютнаяСуммаBalanceCr")) - num(row.get("ВалютнаяСуммаBalanceDr"))
        else:
            party.advance += dr - cr
        party.lines.append({"account": code, "account_name": account_name,
                            "contract": contracts.get(row.get("ExtDimension2"), ""), "dr": dr, "cr": cr,
                            "currency": currencies.get(row.get("Валюта_Key"), ""),
                            "fc": num(row.get("ВалютнаяСуммаBalanceCr")) - num(row.get("ВалютнаяСуммаBalanceDr"))})
    return found


async def read_onec(agent: str = "ap-reconcile") -> dict[str, Any]:
    """What the comparison needs from 1C, read now (GET only): the 60/43 accounts, their balances by
    counterparty and contract, the counterparties (with ИНН), contracts and currencies."""
    from integrations.onec.client import OneCClient

    async with OneCClient(agent=agent) as client:
        chart = (await client.get("ChartOfAccounts_Хозрасчетный", {"$select": "Ref_Key,Code,Description"})).get("value", [])
        chart = [a for a in chart if str(a.get("Code", "")).startswith(("60", "43"))]
        keys = {a["Ref_Key"] for a in chart}
        moment = datetime.now().strftime("%Y-%m-%dT%H:%M:%S")
        rows = (await client.get(f"AccountingRegister_Хозрасчетный/Balance(Period=datetime'{moment}')")).get("value", [])
        return {
            "as_of": moment,
            "chart": chart,
            "rows": [r for r in rows if r.get("Account_Key") in keys],
            "parties": (await client.get("Catalog_Контрагенты", {"$select": "Ref_Key,Description,ИНН"})).get("value", []),
            "contracts": (await client.get("Catalog_ДоговорыКонтрагентов", {"$select": "Ref_Key,Description"})).get("value", []),
            "currencies": (await client.get("Catalog_Валюты", {"$select": "Ref_Key,Code,Description"})).get("value", []),
        }


async def pull_onec(out: Path) -> dict[str, int]:
    """``read_onec`` into a JSON file; returns counts, never amounts."""
    data = await read_onec()
    out.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
    return {"rows": len(data["rows"]), "parties": len(data["parties"])}


# ------------------------------------------------------------------ SAP


@dataclass
class SapParty:
    code: str
    name: str
    inn: str
    owed: float          # so'm: what we owe the supplier (positive)
    owed_lc: float       # SAP's local currency (USD)
    currency: str = ""
    owed_fc: float = 0.0
    som_known: bool = True


def read_table(path: Path) -> list[dict[str, Any]]:
    """An Excel / CSV export as rows of {header: value}."""
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        sheet = load_workbook(path, read_only=True, data_only=True).active
        rows = [list(r) for r in sheet.iter_rows(values_only=True)]
    else:
        text = path.read_text(encoding="utf-8-sig")
        dialect = csv.Sniffer().sniff(text[:4000], delimiters=",;\t")
        rows = list(csv.reader(text.splitlines(), dialect))
    header_at = next((i for i, r in enumerate(rows[:20]) if sum(1 for c in r if str(c or "").strip()) >= 3), 0)
    header = [str(c or "").strip() for c in rows[header_at]]
    return [dict(zip(header, r)) for r in rows[header_at + 1:] if any(str(c or "").strip() for c in r)]


async def pushed_sap_rows(*, close: bool = False) -> tuple[list[dict[str, Any]], str]:
    """The suppliers the gateway last pushed (``supplier_balances``), and when — a READ ONLY read.

    ``close`` closes the database pool afterwards: for the command line only,
    never inside the running service (its pool is shared).
    """
    from integrations.common.db import close_pool, fetch_read_only
    from integrations.common.timeutil import to_local

    try:
        rows = await fetch_read_only(
            "SELECT raw, captured_at FROM v_sap_gateway_latest WHERE tool = 'supplier_balances'"
        )
    finally:
        if close:
            await close_pool()
    raws = [r["raw"] if isinstance(r["raw"], dict) else json.loads(r["raw"]) for r in rows]
    when = max((r["captured_at"] for r in rows), default=None)
    return raws, f"SAP шлюзи, {to_local(when):%d.%m.%Y %H:%M}" if when else "SAP шлюзи (маълумот йўқ)"


def find_columns(header: list[str]) -> dict[str, str]:
    """Our field → the export's column, by header name (exact first, then contains)."""
    lowered = {h.lower().strip(): h for h in header if h}
    found: dict[str, str] = {}
    for key, names in SAP_HEADERS.items():
        exact = next((lowered[n] for n in names if n in lowered and lowered[n] not in found.values()), None)
        if exact is None:
            exact = next((h for low, h in lowered.items() for n in names
                          if n in low and h not in found.values() and len(n) > 3), None)
        if exact:
            found[key] = exact
    return found


def sap_parties(rows: list[dict[str, Any]], sign: float | None = None) -> tuple[dict[str, SapParty], dict[str, Any]]:
    """The export → suppliers, with what we owe each in so'm when SAP gives it (its system currency).

    ``sign`` turns SAP's balance into "we owe = positive": -1 or +1. Left
    None, the majority sign of the balances decides — which proved wrong on
    the real data (09.10.2026: most suppliers hold an advance, stored
    negative), so ``choose_sign`` decides it against 1C instead.
    """
    if not rows:
        return {}, {"columns": {}, "flipped": False}
    columns = find_columns(list(rows[0]))
    if "name" not in columns or not ({"balance_sys", "balance_lc"} & set(columns)):
        raise ValueError(f"SAP export: need a name and a balance column; found {sorted(columns)}")
    kept = [r for r in rows if not ("card_type" in columns and str(r.get(columns["card_type"], "")).strip().upper()
                                    not in ("S", "SUPPLIER", "ПОСТАВЩИК", "ПОСТ."))]
    values = [num(r.get(columns.get("balance_sys") or columns["balance_lc"])) for r in kept]
    # Fallback only (no supplier in both systems): the majority sign is taken as "we owe".
    if sign is None:
        sign = -1.0 if sum(1 for v in values if v < 0) > sum(1 for v in values if v > 0) else 1.0
    flipped = sign < 0
    found: dict[str, SapParty] = {}
    for r in kept:
        name = str(r.get(columns["name"], "") or "").strip()
        if not name:
            continue
        code = str(r.get(columns.get("code", ""), "") or "").strip() or name
        som_known = "balance_sys" in columns
        lc = sign * num(r.get(columns.get("balance_lc", ""), 0))
        owed = sign * num(r.get(columns["balance_sys"])) if som_known else lc
        found[code] = SapParty(code, name, digits(r.get(columns.get("inn", ""), "")), owed, lc,
                               str(r.get(columns.get("currency", ""), "") or "").strip(),
                               sign * num(r.get(columns.get("balance_fc", ""), 0)), som_known)
    return found, {"columns": columns, "flipped": flipped}


# ------------------------------------------------------------------ compare


@dataclass
class Match:
    onec: OneCParty | None
    sap: SapParty | None
    by: str = ""

    @property
    def diff(self) -> float:
        return (self.onec.net if self.onec else 0.0) - (self.sap.owed if self.sap else 0.0)


def choose_sign(onec: dict[str, OneCParty], rows: list[dict[str, Any]]) -> tuple[float | None, dict[str, Any]]:
    """SAP's sign, read from the data: the one under which the suppliers found in both
    systems agree with 1C (smallest total difference).

    On 09.10.2026 the majority guess turned every advance into a debt, so an
    advance equal in both systems showed as twice its size in "difference".
    With no supplier in both systems there is nothing to agree on: None, and
    ``sap_parties`` falls back to the majority guess (the workbook says so).
    """
    best: tuple[float, float, int] | None = None
    for sign in (1.0, -1.0):
        sap, _ = sap_parties(rows, sign)
        pairs = [m for m in match(onec, sap) if m.onec and m.sap]
        total = sum(abs(m.diff) for m in pairs)
        if best is None or total < best[1]:
            best = (sign, total, len(pairs))
    assert best is not None
    if best[2] == 0:
        return None, {"sign_by_1c": False, "matched": 0}
    return best[0], {"sign_by_1c": True, "matched": best[2]}


def tolerance(amount: float) -> float:
    return max(TOLERANCE_SOM, abs(amount) * TOLERANCE_SHARE)


def match(onec: dict[str, OneCParty], sap: dict[str, SapParty]) -> list[Match]:
    """Pair counterparties: by ИНН first, then by the cleaned-up name; the rest stay alone."""
    left = dict(onec)
    pairs: list[Match] = []
    by_inn = {p.inn: k for k, p in left.items() if len(p.inn) >= 9}
    by_name: dict[str, str] = {}
    for k, p in left.items():
        by_name.setdefault(norm_name(p.name), k)
    for s in sorted(sap.values(), key=lambda s: -abs(s.owed)):
        key, how = None, ""
        if len(s.inn) >= 9 and by_inn.get(s.inn) in left:
            key, how = by_inn[s.inn], "ИНН"
        elif by_name.get(norm_name(s.name)) in left:
            key, how = by_name[norm_name(s.name)], "ном"
        pairs.append(Match(left.pop(key) if key else None, s, how))
    pairs += [Match(p, None) for p in left.values()]
    return pairs


def reason(m: Match) -> str:
    """The likely reason for a difference, in Uzbek Cyrillic — a hint for the accountant, not a verdict."""
    o, s = m.onec, m.sap
    if o is None:
        return "Фақат SAP'да: 1C'да бу контрагент бўйича қолдиқ йўқ — ҳужжат 1C'га киритилмаган ёки ёпилган"
    if s is None:
        return "Фақат 1C'да: SAP'да бу етказиб берувчи бўйича қолдиқ йўқ — ҳужжат SAP'га киритилмаган ёки контрагент бошқа номда"
    d = m.diff
    if abs(d) <= tolerance(max(abs(o.net), abs(s.owed))):
        return "Мос"
    if o.advance and abs(abs(d) - o.advance) <= tolerance(o.advance):
        return "Аванс фарқи: 1C'да тўланган аванс (4310) ҳисобга олинган, SAP'да эмас (ёки аксинча)"
    if o.unoffset and abs(abs(d) - o.advance) > tolerance(o.advance):
        return UNOFFSET + " — SAP битта соф қолдиқ кўрсатади"
    if (o.net > 0) != (s.owed > 0) and o.net and s.owed:
        return "Йўналиш ҳар хил: бир тизимда қарзмиз, бошқасида ортиқча тўлов/аванс — тўлов бир тизимда киритилмаган"
    if not s.som_known:
        return "SAP суммаси сўмда эмас (USD): курс фарқи бўлиши мумкин — SAP'дан сўмдаги қолдиқ (BalanceSys) керак"
    if o.payable_fc and abs(d) <= 0.05 * max(abs(o.net), abs(s.owed)):
        return "Валюта курси: валютадаги қарз (6015) икки тизимда турли курсда қайта баҳоланган"
    return "Сумма фарқли: бир тизимда ҳужжат ёки тўлов йўқ ёки суммаси бошқа — акт-сверка қилиш керак"


# ------------------------------------------------------------------ workbook

LANGS = ("uz", "ru")
# The SAP report a bookkeeper opens to see the same SAP figure (2026-10-10:
# SAP's own column names differ from the workbook's, so each sheet explains it).
_SAP_REPORT = {
    "uz": "SAP'даги «Кредиторская задолженность» ҳисоботи (Валюта = Система), «К оплате» устунидаги етказиб берувчи жами",
    "ru": "отчёт SAP «Кредиторская задолженность» (Валюта = Система), столбец «К оплате» — итог по поставщику",
}
# Every text of the workbook, per language. Uzbek is the source; reason() returns it, and REASONS_RU translates.
TEXT: dict[str, dict[str, Any]] = {
    "uz": {
        "sheets": ("Солиштириш", "Фақат 1C", "Фақат SAP", "1C тафсилот", "Изоҳ"),
        "compare": ["Контрагент (1C)", "Контрагент (SAP)", "ИНН", "Боғланди", "1C: қарз (6010+6015), сўм",
                    "1C: аванс (4310+4315), сўм", "1C: нетто, сўм", "SAP: қарз, сўм", "SAP: қарз, USD",
                    "Фарқ (1C − SAP), сўм", "Эҳтимолий сабаб"],
        "only_1c": ["Контрагент", "ИНН", "Қарз (6010+6015), сўм", "Аванс (4310+4315), сўм", "Нетто, сўм",
                    "Валютада (6015)", "Эҳтимолий сабаб", "1C ичида"],
        "only_sap": ["Код", "Контрагент", "ИНН", "Қарз, сўм", "Қарз, USD", "Валюта", "Эҳтимолий сабаб"],
        "detail": ["Контрагент", "Ҳисобварақ", "Номи", "Шартнома", "Дт қолдиқ, сўм", "Кт қолдиқ, сўм", "Валютада", "Валюта"],
        "by": {"ИНН": "ИНН", "ном": "ном"},
        "no_sap": "SAP маълумоти ҳали келмаган — солиштирилмади",
        "legend": "УСТУНЛАР ИЗОҲИ",
        "legend_head": ("Устун", "Нимани билдиради ва SAP / 1C'да қаерда"),
        "notes": (
            "1C қолдиқлари: {as_of} ҳолатига (OData, фақат ўқиш). SAP: {sap}.",
            "1C: қарз = 6010 + 6015 ҳисобварақларининг кредит қолдиғи; аванс = 4310 + 4315 дебет қолдиғи; нетто = қарз − аванс.",
            "SAP: етказиб берувчилар қолдиғи (OCRD, тури S); сўмда — SAP'нинг тизим валютаси (BalanceSys), бўлмаса USD.",
            "Боғлаш: аввал ИНН бўйича, кейин тозаланган ном бўйича (ООО, МЧЖ, қўштирноқлар олиб ташланган).",
            "Фарқ кичик бўлса (≤ {tol} сўм ёки 0,5 %) «Мос» деб олинади.",
            "Сабаб — тахмин, ҳукм эмас: ҳар бир фарқ бухгалтерия томонидан акт-сверка билан текширилади.",
            "SAP'даги қолдиқ бугунги ҳолатга. SAP ҳисоботида «Дата расчета» бошқа кун бўлса, ўша кундан кейин киритилганлар фарқ қилади.",
        ),
        "columns": "SAP устунлари топилди: ",
        "flipped": "; SAP қолдиғининг ишораси алмаштирилди",
        "sign_1c": "; ишора 1C билан мос келишига қараб танланди",
        "sign_guess": "; иккала тизимда ҳам бор етказиб берувчи йўқ — ишора кўпчилик қолдиққа қараб олинди",
        "explain": {
            "compare": (
                "1C'даги «Контрагенты» маълумотномасидаги номи.",
                "SAP'даги номи: «Название поставщика» устуни (коди — «Код поставщика»).",
                "Контрагентнинг ИНН (СТИР). SAP'да — ҳамкор карточкасидаги «Федеральный налоговый номер» (LicTradNum).",
                "Жуфтлик қандай топилди: «ИНН» — бир хил ИНН бўйича; «ном» — ИНН йўқ, номи мос келди (ООО, МЧЖ, қўштирноқсиз).",
                "1C бўйича БИЗ етказиб берувчига қанча ҚАРЗМИЗ: 6010 (сўм) ва 6015 (валюта, сўмга ҳисобланган) кредит қолдиғи.",
                "1C бўйича ОЛДИНДАН ТЎЛАГАНИМИЗ (берилган аванс): 4310 ва 4315 дебет қолдиғи.",
                "1C жами: қарз минус аванс. Мусбат — биз қарздормиз; манфий — етказиб берувчи бизга қарздор (ортиқча тўлов/аванс).",
                "SAP жами, сўмда: етказиб берувчининг тизим валютасидаги сальдоси. Бу {sap}. Мусбат — биз қарздормиз; манфий — аванс/ортиқча тўлов.",
                "Худди шу сальдо SAP'нинг асосий валютасида (USD) — ўша ҳисоботда «Валюта = Основная».",
                "1C нетто минус SAP. 0 — мос. Мусбат — 1C'да қарз SAP'дагидан кўп; манфий — SAP'да кўп.",
                "Нега мос келмагани ҳақида тахмин. Бу хулоса эмас — ҳар бир фарқ акт-сверка билан текширилади.",
            ),
            "only_1c": (
                "1C'да қолдиғи бор, SAP'да жуфти топилмаган етказиб берувчи (на ИНН, на ном бўйича).",
                "1C'даги ИНН.",
                "1C бўйича қарзимиз (6010 + 6015 кредит).",
                "1C бўйича берилган аванс (4310 + 4315 дебет).",
                "Қарз минус аванс. Мусбат — биз қарздормиз; манфий — бизга қарздор.",
                "6015 бўйича валютадаги қарз ўз валютасида (USD, EUR …).",
                "Нега SAP'да йўқ: ҳужжат SAP'га киритилмаган ёки SAP'да бошқа ном билан / ИНН'сиз — SAP ҳисоботидан номи бўйича қидиринг.",
                "Белгиланган бўлса — 1C'да бир вақтда қарз ҳам, аванс ҳам бор: аванс зачет қилинмаган бўлиши мумкин.",
            ),
            "only_sap": (
                "SAP'даги «Код поставщика».",
                "SAP'даги «Название поставщика».",
                "SAP'даги етказиб берувчи карточкасидан ИНН (кўпинча бўш — унда жуфти номи бўйича қидирилади).",
                "SAP'даги сальдо, сўмда = {sap}. Мусбат — биз қарздормиз; манфий — аванс.",
                "Худди шу сальдо USD'да (Валюта = Основная).",
                "SAP'даги етказиб берувчи карточкасининг валютаси.",
                "Нега 1C'да йўқ: ҳужжат 1C'га киритилмаган ёки у ерда ёпилган. Кўпинча ходимлар ва карталар — 1C'да улар бошқа ҳисобварақларда.",
            ),
            "detail": (
                "1C'даги контрагент.",
                "1C ҳисобварағи: 6010 / 6015 — етказиб берувчилар билан ҳисоб-китоб (сўм / валюта); 4310 / 4315 — берилган аванслар.",
                "1C ҳисобвараклар режасидаги номи.",
                "1C'даги контрагент шартномаси (ҳисобварақнинг иккинчи аналитикаси).",
                "Дебет қолдиқ: 43xx учун — берилган аванс; 60xx учун — ортиқча тўлов.",
                "Кредит қолдиқ: 60xx учун — етказиб берувчига қарзимиз.",
                "Шартнома валютасидаги қолдиқ (6015 учун).",
                "Шу қолдиқнинг валютаси.",
            ),
        },
    },
    "ru": {
        "sheets": ("Сверка", "Только в 1С", "Только в SAP", "1С детализация", "Пояснения"),
        "compare": ["Контрагент (1С)", "Контрагент (SAP)", "ИНН", "Сопоставлено по", "1С: долг (6010+6015), сум",
                    "1С: аванс (4310+4315), сум", "1С: нетто, сум", "SAP: долг, сум", "SAP: долг, USD",
                    "Разница (1С − SAP), сум", "Вероятная причина"],
        "only_1c": ["Контрагент", "ИНН", "Долг (6010+6015), сум", "Аванс (4310+4315), сум", "Нетто, сум",
                    "В валюте (6015)", "Вероятная причина", "Внутри 1С"],
        "only_sap": ["Код", "Контрагент", "ИНН", "Долг, сум", "Долг, USD", "Валюта", "Вероятная причина"],
        "detail": ["Контрагент", "Счёт", "Наименование", "Договор", "Остаток Дт, сум", "Остаток Кт, сум", "В валюте", "Валюта"],
        "by": {"ИНН": "ИНН", "ном": "наименование"},
        "no_sap": "Данных SAP ещё нет — не сверено",
        "legend": "ПОЯСНЕНИЕ К СТОЛБЦАМ",
        "legend_head": ("Столбец", "Что означает и где найти в SAP / 1С"),
        "notes": (
            "Остатки 1С: на {as_of} (OData, только чтение). SAP: {sap}.",
            "1С: долг = кредитовый остаток счетов 6010 + 6015; аванс = дебетовый остаток 4310 + 4315; нетто = долг − аванс.",
            "SAP: остатки поставщиков (OCRD, тип S); в сумах — системная валюта SAP (BalanceSys), иначе USD.",
            "Сопоставление: сначала по ИНН, затем по очищенному наименованию (без ООО, МЧЖ, кавычек).",
            "Если разница мала (≤ {tol} сум или 0,5 %), считается «Совпадает».",
            "Причина — предположение, а не вывод: каждое расхождение бухгалтерия проверяет актом сверки.",
            "Остаток SAP — на сегодня. Если в отчёте SAP «Дата расчета» другая, суммы отличаются на то, что введено после неё.",
        ),
        "columns": "Найденные столбцы SAP: ",
        "flipped": "; знак остатка SAP изменён",
        "sign_1c": "; знак выбран по совпадению с 1С",
        "sign_guess": "; нет поставщиков в обеих системах — знак взят по большинству остатков",
        "explain": {
            "compare": (
                "Наименование контрагента в справочнике «Контрагенты» 1С.",
                "Наименование в SAP: столбец «Название поставщика» (код — «Код поставщика»).",
                "ИНН (СТИР) контрагента. В SAP — «Федеральный налоговый номер» (LicTradNum) в карточке делового партнёра.",
                "Как найдена пара: «ИНН» — по одинаковому ИНН; «наименование» — ИНН нет, совпало название (без ООО, МЧЖ, кавычек).",
                "Сколько МЫ ДОЛЖНЫ поставщику по 1С: кредитовый остаток 6010 (сум) и 6015 (валюта, в пересчёте на сумы).",
                "Сколько мы ЗАПЛАТИЛИ ВПЕРЁД (выданный аванс) по 1С: дебетовый остаток 4310 и 4315.",
                "Итог по 1С: долг минус аванс. Плюс — мы должны; минус — поставщик должен нам (переплата/аванс).",
                "Итог по SAP в сумах: сальдо поставщика в системной валюте. Это {sap}. Плюс — мы должны; минус — аванс/переплата.",
                "То же сальдо в основной валюте SAP (USD) — «Валюта = Основная» в том же отчёте.",
                "1С нетто минус SAP. 0 — совпадает. Плюс — в 1С долг больше, чем в SAP; минус — в SAP больше.",
                "Подсказка, почему суммы не совпали. Это не вывод — каждое расхождение проверяется актом сверки.",
            ),
            "only_1c": (
                "Поставщик, у которого остаток есть в 1С, а в SAP пары не нашлось (ни по ИНН, ни по названию).",
                "ИНН из 1С.",
                "Сколько мы должны по 1С (кредит 6010 + 6015).",
                "Выданный аванс по 1С (дебет 4310 + 4315).",
                "Долг минус аванс. Плюс — мы должны; минус — должны нам.",
                "Валютный долг по 6015 в своей валюте (USD, EUR …).",
                "Почему нет в SAP: документ не введён в SAP, или в SAP поставщик под другим названием / без ИНН — найдите его в отчёте SAP по названию.",
                "Отмечено, если в 1С у поставщика одновременно и долг, и аванс — возможно, аванс не зачтён.",
            ),
            "only_sap": (
                "«Код поставщика» в SAP.",
                "«Название поставщика» в SAP.",
                "ИНН из карточки поставщика в SAP (часто пусто — тогда пару ищут по названию).",
                "Сальдо поставщика в SAP в сумах = {sap}. Плюс — мы должны; минус — аванс.",
                "То же сальдо в USD (Валюта = Основная).",
                "Валюта карточки поставщика в SAP.",
                "Почему нет в 1С: документ не введён в 1С или закрыт там. Часто это сотрудники и карты — в 1С они на других счетах.",
            ),
            "detail": (
                "Контрагент из 1С.",
                "Счёт 1С: 6010 / 6015 — расчёты с поставщиками (сум / валюта); 4310 / 4315 — выданные авансы.",
                "Название счёта в плане счетов 1С.",
                "Договор контрагента в 1С (вторая аналитика счёта).",
                "Дебетовый остаток: для 43xx — выданный аванс; для 60xx — переплата.",
                "Кредитовый остаток: для 60xx — наш долг поставщику.",
                "Остаток в валюте договора (для 6015).",
                "Валюта этого остатка.",
            ),
        },
    },
}
REASONS_RU = {
    "Мос": "Совпадает",
    "Фақат SAP'да: 1C'да бу контрагент бўйича қолдиқ йўқ — ҳужжат 1C'га киритилмаган ёки ёпилган":
        "Только в SAP: в 1С нет остатка по этому контрагенту — документ не введён в 1С или закрыт",
    "Фақат 1C'да: SAP'да бу етказиб берувчи бўйича қолдиқ йўқ — ҳужжат SAP'га киритилмаган ёки контрагент бошқа номда":
        "Только в 1С: в SAP нет остатка по этому поставщику — документ не введён в SAP или контрагент под другим названием",
    "Аванс фарқи: 1C'да тўланган аванс (4310) ҳисобга олинган, SAP'да эмас (ёки аксинча)":
        "Разница в авансе: выданный аванс (4310) учтён в 1С, но не в SAP (или наоборот)",
    UNOFFSET + " — SAP битта соф қолдиқ кўрсатади":
        "В 1С одновременно есть и долг, и аванс — возможно, аванс не зачтён; SAP показывает одно сальдо",
    UNOFFSET: "В 1С одновременно есть и долг, и аванс — возможно, аванс не зачтён",
    "Йўналиш ҳар хил: бир тизимда қарзмиз, бошқасида ортиқча тўлов/аванс — тўлов бир тизимда киритилмаган":
        "Разное направление: в одной системе мы должны, в другой — переплата/аванс; оплата введена только в одной системе",
    "SAP суммаси сўмда эмас (USD): курс фарқи бўлиши мумкин — SAP'дан сўмдаги қолдиқ (BalanceSys) керак":
        "Сумма SAP не в сумах (USD): возможна курсовая разница — нужен остаток SAP в сумах (BalanceSys)",
    "Валюта курси: валютадаги қарз (6015) икки тизимда турли курсда қайта баҳоланган":
        "Курс валюты: валютный долг (6015) переоценён в двух системах по разным курсам",
    "Сумма фарқли: бир тизимда ҳужжат ёки тўлов йўқ ёки суммаси бошқа — акт-сверка қилиш керак":
        "Суммы различаются: в одной системе нет документа или оплаты, либо сумма другая — нужен акт сверки",
}


def say(text: str, lang: str) -> str:
    """A reason (or other Uzbek text of the workbook) in the workbook's language."""
    return REASONS_RU.get(text, text) if lang == "ru" else text


def write_workbook(pairs: list[Match], onec: dict[str, OneCParty], meta: dict[str, Any],
                   out: Path | IO[bytes], lang: str = "uz") -> dict[str, int]:
    """The Excel file for the Director and accounting, in Uzbek or Russian; returns counts (never amounts).

    Under each table: what every column means and where the same figure is
    in SAP / 1C (2026-10-10 — the accountant asked; SAP's own report names
    its columns differently).
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    t = TEXT[lang if lang in LANGS else "uz"]
    sap_report = _SAP_REPORT[lang if lang in LANGS else "uz"]
    wb = Workbook()
    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="E8EEF4")
    bad_fill = PatternFill("solid", fgColor="FDECEC")
    wrap = Alignment(wrap_text=True, vertical="top")
    money = '#,##0;[Red]-#,##0'

    def sheet(title: str, header: list[str], widths: list[int]):
        ws = wb.create_sheet(title)
        ws.append(header)
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = w
            ws.cell(1, i).font, ws.cell(1, i).fill = bold, head_fill
            ws.cell(1, i).alignment = wrap
        ws.freeze_panes = "A2"
        return ws

    def legend(ws, header: list[str], key: str) -> None:
        """Every column explained, two rows under the table."""
        row = ws.max_row + 3
        ws.cell(row, 1, t["legend"]).font = bold
        row += 1
        for i, text in enumerate(t["legend_head"], start=1):
            ws.cell(row, i, text).font = bold
            ws.cell(row, i).fill = head_fill
        for name, text in zip(header, t["explain"][key]):
            row += 1
            ws.cell(row, 1, name).font = bold
            ws.cell(row, 2, text.format(sap=sap_report)).alignment = wrap
            ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=max(2, min(len(header), 8)))
            ws.row_dimensions[row].height = 32

    wb.remove(wb.active)
    names = t["sheets"]
    both = [m for m in pairs if m.onec and m.sap]
    ws = sheet(names[0], t["compare"], [34, 34, 13, 10, 18, 18, 16, 16, 14, 16, 60])
    for m in sorted(both, key=lambda m: -abs(m.diff)):
        why = reason(m)
        ws.append([m.onec.name, m.sap.name, m.onec.inn or m.sap.inn, t["by"].get(m.by, m.by), m.onec.payable,
                   m.onec.advance, m.onec.net, m.sap.owed, m.sap.owed_lc, m.diff, say(why, lang)])
        if why != "Мос":
            for cell in ws[ws.max_row]:
                cell.fill = bad_fill
    for col in "EFGHIJ":
        for cell in ws[col][1:]:
            cell.number_format = money
    legend(ws, t["compare"], "compare")

    only_1c = [m.onec for m in pairs if m.onec and not m.sap and abs(m.onec.net) > TOLERANCE_SOM]
    ws = sheet(names[1], t["only_1c"], [40, 13, 18, 18, 16, 22, 60, 44])
    for p in sorted(only_1c, key=lambda p: -abs(p.net)):
        ws.append([p.name, p.inn, p.payable, p.advance, p.net,
                   "; ".join(f"{v:,.2f} {c}" for c, v in p.payable_fc.items()),
                   say(reason(Match(p, None)), lang) if meta.get("sap_rows") else t["no_sap"],
                   say(UNOFFSET, lang) if p.unoffset else ""])
    for col in "CDE":
        for cell in ws[col][1:]:
            cell.number_format = money
    legend(ws, t["only_1c"], "only_1c")

    only_sap = [m.sap for m in pairs if m.sap and not m.onec and abs(m.sap.owed) > TOLERANCE_SOM]
    ws = sheet(names[2], t["only_sap"], [14, 40, 13, 16, 14, 10, 60])
    for s in sorted(only_sap, key=lambda s: -abs(s.owed)):
        ws.append([s.code, s.name, s.inn, s.owed, s.owed_lc, s.currency, say(reason(Match(None, s)), lang)])
    for col in "DE":
        for cell in ws[col][1:]:
            cell.number_format = money
    legend(ws, t["only_sap"], "only_sap")

    ws = sheet(names[3], t["detail"], [36, 10, 36, 36, 16, 16, 14, 8])
    for p in sorted(onec.values(), key=lambda p: p.name):
        for line in p.lines:
            ws.append([p.name, line["account"], line["account_name"], line["contract"], line["dr"], line["cr"],
                       line["fc"] or None, line["currency"] if line["fc"] else ""])
    for col in "EF":
        for cell in ws[col][1:]:
            cell.number_format = money
    legend(ws, t["detail"], "detail")

    ws = sheet(names[4], [names[4]], [120])
    notes = [n.format(as_of=meta.get("onec_as_of", ""), sap=meta.get("sap_file", ""), tol=f"{TOLERANCE_SOM:,.0f}")
             for n in t["notes"]]
    notes.append(t["columns"] + ", ".join(f"{k}={v}" for k, v in meta.get("sap_columns", {}).items())
                 + (t["flipped"] if meta.get("flipped") else "")
                 + (t["sign_1c"] if meta.get("sign_by_1c") else t["sign_guess"]))
    for line in notes:
        ws.append([line])
    wb.save(out)
    return {"matched": len(both), "differ": sum(1 for m in both if reason(m) != "Мос"),
            "only_1c": len(only_1c), "only_sap": len(only_sap)}


def workbook_bytes(result: Comparison, lang: str = "uz") -> bytes:
    """``write_workbook`` in memory — what the bot sends as an .xlsx."""
    buffer = io.BytesIO()
    write_workbook(result.pairs, result.onec, result.meta, buffer, lang)
    return buffer.getvalue()


def workbook_name(day: Any, lang: str = "uz") -> str:
    """The file's name as the bot sends it."""
    return f"kreditorskaya-1C-SAP-{day:%Y-%m-%d}-RU.xlsx" if lang == "ru" else f"kreditorlik-1C-SAP-{day:%Y-%m-%d}.xlsx"


# ------------------------------------------------------------------ in one go


@dataclass
class Comparison:
    pairs: list[Match]
    onec: dict[str, OneCParty]
    meta: dict[str, Any]

    @property
    def both(self) -> list[Match]:
        return [m for m in self.pairs if m.onec and m.sap]


def compare(onec_data: dict[str, Any], sap_rows: list[dict[str, Any]], sap_source: str) -> Comparison:
    """1C's data (``read_onec``) against SAP's supplier rows: pairs, with SAP's sign chosen by 1C."""
    onec = onec_parties(onec_data)
    if sap_rows:
        sign, how = choose_sign(onec, sap_rows)
        sap, info = sap_parties(sap_rows, sign)
        info["sign_by_1c"] = how["sign_by_1c"]
    else:
        sap, info = {}, {"columns": {}, "flipped": False}
    meta = {"onec_as_of": onec_data.get("as_of", ""), "sap_file": sap_source, "sap_columns": info["columns"],
            "flipped": info["flipped"], "sign_by_1c": info.get("sign_by_1c", False), "sap_rows": len(sap)}
    return Comparison(match(onec, sap), onec, meta)


def _som(value: float) -> str:
    return f"{value:,.0f}".replace(",", " ") + " сўм"


def summary(result: Comparison, top: int = 10) -> str:
    """The comparison as a short text for the analyst (it words the answer): totals, counts, the biggest gaps."""
    pairs, both = result.pairs, result.both
    meta = result.meta
    lines = [f"Supplier debt (кредиторлик), 1C vs SAP. 1C balances as of {meta.get('onec_as_of', '')[:16].replace('T', ' ')}; "
             f"SAP: {meta.get('sap_file', '')}. Amounts in so'm; positive = we owe the supplier, "
             "negative = we paid an advance / overpaid."]
    if not meta.get("sap_rows"):
        lines.append("SAP has no supplier balances yet (the gateway's get_supplier_balances hasn't pushed) — "
                     "only 1C's side is known.")
    onec_net = sum(p.net for p in result.onec.values())
    onec_payable = sum(p.payable for p in result.onec.values())
    onec_advance = sum(p.advance for p in result.onec.values())
    sap_owed = sum(m.sap.owed for m in pairs if m.sap)
    lines.append(f"TOTAL 1C: owed {_som(onec_payable)} (6010+6015), advances paid {_som(onec_advance)} (4310+4315), "
                 f"net {_som(onec_net)}.")
    if meta.get("sap_rows"):
        lines.append(f"TOTAL SAP: net {_som(sap_owed)} ({meta['sap_rows']} suppliers with a balance). "
                     f"Difference 1C − SAP: {_som(onec_net - sap_owed)}.")
    differ = [m for m in both if reason(m) != "Мос"]
    only_1c = [m.onec for m in pairs if m.onec and not m.sap and abs(m.onec.net) > TOLERANCE_SOM]
    only_sap = [m.sap for m in pairs if m.sap and not m.onec and abs(m.sap.owed) > TOLERANCE_SOM]
    lines.append(f"Suppliers in both systems: {len(both)} — agree {len(both) - len(differ)}, differ {len(differ)}. "
                 f"Only in 1C: {len(only_1c)}. Only in SAP: {len(only_sap)}.")
    if differ:
        lines.append("Biggest differences (1C net / SAP / difference — likely reason, a hint for accounting):")
        for m in sorted(differ, key=lambda m: -abs(m.diff))[:top]:
            lines.append(f"  - {m.onec.name}: 1C {_som(m.onec.net)} / SAP {_som(m.sap.owed)} / {_som(m.diff)} — {reason(m)}")
    if only_1c:
        lines.append("Biggest only in 1C (SAP may hold them under another name):")
        lines += [f"  - {p.name}: {_som(p.net)}" for p in sorted(only_1c, key=lambda p: -abs(p.net))[:top // 2]]
    if only_sap:
        lines.append("Biggest only in SAP:")
        lines += [f"  - {s.name}: {_som(s.owed)}" for s in sorted(only_sap, key=lambda s: -abs(s.owed))[:top // 2]]
    if not meta.get("sign_by_1c") and meta.get("sap_rows"):
        lines.append("Note: no supplier is in both systems, so SAP's sign couldn't be checked against 1C.")
    return "\n".join(lines)
