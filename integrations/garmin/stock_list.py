"""Garmin stock — every item and how many are left (2026-10-27, Murod aka).

«Барча товарлар остаткаси керак, фақат Garmin'га тегишли SAP'даги маълумот.» /eski stays (old, unsold goods); this
is the whole Garmin stock: each product's name and how many pieces are left, per Garmin warehouse and in total. The
same day the owner cut it down to just that: «it is not your responsibility to say what to buy — OPS Manager just
says how much is left and the name of the product». So: no marks, no sales pace, no advice on what to order.

/qoldiq (the Garmin manager, the Director) sends a short summary and the .xlsx; asked in words, the analyst's
garmin_stock_list tool gives the same and attaches the file. Quantity only — no money. Display dummies and the
faulty-goods warehouse («Faulty Product Garmin») left out.
"""

from __future__ import annotations

import io
import json
import uuid
from dataclasses import dataclass, field
from typing import Any

from integrations.common.config import settings
from integrations.common.db import fetch_all
from integrations.common.logging_setup import setup_logging
from integrations.common.timeutil import to_local, today_local

AGENT = "garmin-stock"
log = setup_logging(AGENT)

COMMANDS = ("/qoldiq", "/қолдиқ", "/ostatka")
FAULTY = ("faulty", "брак", "brak", "nosoz", "носоз", "defect")  # broken goods aren't stock to sell (27.10)


@dataclass
class Item:
    code: str
    name: str
    per_wh: dict[str, float] = field(default_factory=dict)

    @property
    def total(self) -> float:
        return sum(self.per_wh.values())


def sellable(codes: frozenset[str], warehouses: list[dict[str, Any]]) -> frozenset[str]:
    """The Garmin warehouses without the faulty-goods ones («Faulty Product Garmin», 06)."""
    names = {str(w.get("WhsCode") or "").strip(): str(w.get("WhsName") or "").lower() for w in warehouses}
    return frozenset(c for c in codes if not any(word in names.get(c, "") for word in FAULTY))


def _num(value: Any) -> float:
    try:
        return float(value or 0)
    except (TypeError, ValueError):
        return 0.0


def build(stock: list[dict[str, Any]], codes: frozenset[str]) -> list[Item]:
    """Every Garmin item with pieces left, by name."""
    from integrations.org_bot.old_stock import _DUMMY  # noqa: PLC2701 — the one definition of a display sample

    items: dict[str, Item] = {}
    for r in stock:
        whs, code, name = str(r.get("WhsCode") or "").strip(), str(r.get("ItemCode") or ""), str(r.get("ItemName") or "")
        qty = _num(r.get("OnHand"))
        if whs in codes and qty > 0 and not (_DUMMY.search(name) or _DUMMY.search(code)):
            item = items.setdefault(code, Item(code, name or code))
            item.per_wh[whs] = item.per_wh.get(whs, 0.0) + qty
    return sorted(items.values(), key=lambda i: i.name.lower())


def _q(value: float) -> str:
    return f"{value:,.0f}".replace(",", " ") if value == int(value) else f"{value:,.1f}".replace(",", " ")


def summary(items: list[Item], warehouses: dict[str, str], as_of: Any = None) -> str:
    """The chat message: how many pieces and products are left, per warehouse — nothing more."""
    from integrations.telegram.bot import escape

    when = f" · SAP {to_local(as_of):%d.%m %H:%M}" if as_of else ""
    lines = [f"📦 <b>Garmin қолдиғи</b>{when}",
             f"Жами: <b>{_q(sum(i.total for i in items))} дона</b>, {len(items)} хил товар"]
    per_wh: dict[str, float] = {}
    for i in items:
        for c, q in i.per_wh.items():
            per_wh[c] = per_wh.get(c, 0.0) + q
    lines += [f"• {escape(warehouses.get(c) or c)} — {_q(q)} дона" for c, q in sorted(per_wh.items(), key=lambda x: -x[1])]
    lines.append("\nҲар бир товар ва унинг қолдиғи — Excel файлда.")
    return "\n".join(lines)


def workbook_bytes(items: list[Item], warehouses: dict[str, str], as_of: Any = None) -> bytes:
    """The whole list as .xlsx: product, code, a column per Garmin warehouse, the total; filters on, header frozen."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    used = [c for c in warehouses if any(c in i.per_wh for i in items)]
    wb = Workbook()
    ws = wb.active
    ws.title = "Қолдиқ"
    ws.append([f"Garmin қолдиғи · {today_local():%d.%m.%Y}"
               + (f" · SAP {to_local(as_of):%d.%m.%Y %H:%M} ҳолатида" if as_of else "")])
    ws["A1"].font = Font(bold=True, size=13)
    head = ["Товар", "Код", *[warehouses[c] or c for c in used], "Жами қолдиқ"]
    ws.append(head)
    for cell in ws[2]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for i in items:
        ws.append([i.name, i.code, *[i.per_wh.get(c, 0) for c in used], i.total])
    ws.freeze_panes = "C3"
    ws.auto_filter.ref = f"A2:{ws.cell(row=2, column=len(head)).column_letter}{ws.max_row}"
    for n, w in enumerate([50, 16, *[12] * len(used), 12], start=1):
        ws.column_dimensions[ws.cell(row=2, column=n).column_letter].width = w
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def file_name() -> str:
    return f"garmin-qoldiq-{today_local():%Y-%m-%d}.xlsx"


async def load() -> tuple[list[Item], dict[str, str], Any]:
    """(items, Garmin warehouse code → name, when SAP's stock arrived) from the latest push."""
    from integrations.org_bot import old_stock

    async def kind(tool: str) -> list[dict[str, Any]]:
        rows = await fetch_all("SELECT raw FROM v_sap_gateway_latest WHERE tool = %s", (tool,))
        return [r["raw"] if isinstance(r["raw"], dict) else json.loads(r["raw"] or "{}") for r in rows]

    whs = await kind("warehouses")
    codes = sellable(old_stock.garmin_codes(whs), whs)
    names = {str(w.get("WhsCode") or "").strip(): str(w.get("WhsName") or "") for w in whs}
    at = await fetch_all("SELECT max(captured_at) AS at FROM v_sap_gateway_latest WHERE tool = 'stock_items'")
    return build(await kind("stock_items"), codes), {c: names.get(c, c) for c in sorted(codes)}, at[0]["at"] if at else None


async def send(chat_id: int, run_id: uuid.UUID) -> str:
    """/qoldiq: the summary and the .xlsx to one chat."""
    from integrations.telegram.bot import TelegramBot

    items, warehouses, as_of = await load()
    async with TelegramBot(agent=AGENT, run_id=run_id,
                           bot_token=settings.ops_manager_bot_telegram_bot_token.get_secret_value()) as bot:
        if not items:
            await bot.send_message("📦 Garmin омборларида SAP'дан қолдиқ маълумоти келмаган.", chat_id=str(chat_id))
            return "garmin_stock_empty"
        await bot.send_message(summary(items, warehouses, as_of), chat_id=str(chat_id))
        await bot.send_file(workbook_bytes(items, warehouses, as_of), file_name(), str(chat_id))
    return "garmin_stock_sent"
