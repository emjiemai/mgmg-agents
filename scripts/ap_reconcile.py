"""Supplier debt (кредиторлик): 1C against SAP B1, supplier by supplier — read-only.

The Director's task of 08.10.2026: "1C билан SAP B1даги кредиторлик
қарздорлигини солиштириб чиқинг — қайси контрагентлар ва суммаларда фарқ
бор, сабаби нима". Done under his written temporary permission of 09.10
(``docs/access-review-2026-10-07.md``, section 4).

Reads, never writes:
  * 1C — over OData with the GET-only ``OneCClient``: balances of 6010 / 6015
    (payables to suppliers, so'm and currency) and 4310 / 4315 (advances paid
    to suppliers), by counterparty and contract, with each counterparty's ИНН;
  * SAP — a file exported from SAP B1 (the gateway is on its own computer and
    doesn't publish suppliers yet): the suppliers' balances (OCRD, CardType
    'S'), as Excel or CSV. Columns are found by their header (SAP's English or
    Russian export names, or the field names).

Writes one Excel workbook (``--out``) for the Director and accounting: the
comparison, what is only in 1C, what is only in SAP, the 1C detail by
contract, and how to read the reasons. Nothing goes into the database, 1C,
SAP or Telegram, and no figure is printed or logged — only counts.

Run:
    python scripts/ap_reconcile.py --onec-json 1c.json --sap sap.xlsx --out solishtirish.xlsx
    python scripts/ap_reconcile.py --pull-1c 1c.json      # read 1C now (needs ONEC_* set)
"""

from __future__ import annotations

import argparse
import asyncio
import csv
import json
import re
import sys
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

PAYABLE = ("6010", "6015")    # расчеты с поставщиками и подрядчиками (сўм, валюта)
ADVANCE = ("4310", "4315")    # авансы выданные (supplier advances)
TOLERANCE_SOM = 1000.0
UNOFFSET = "1C'да бир вақтда қарз ҳам, аванс ҳам бор — аванс зачет қилинмаган бўлиши мумкин"
TOLERANCE_SHARE = 0.005       # 0.5 %

# SAP export headers -> our field (lower-cased; English UI, Russian UI, field names)
SAP_HEADERS = {
    "code": ("cardcode", "bp code", "код бп", "код партнера", "код делового партнера", "код"),
    "name": ("cardname", "bp name", "имя бп", "наименование бп", "наименование", "название", "имя"),
    "inn": ("lictradnum", "federal tax id", "инн", "tin", "stir", "стир", "налоговый номер"),
    "balance_sys": ("balancesys", "balance (sc)", "сальдо (св)", "баланс (св)", "account balance (sc)"),
    "balance_lc": ("balance", "account balance", "сальдо", "баланс", "сальдо счета"),
    "balance_fc": ("balancefc", "balance (fc)", "сальдо (ин. вал.)", "баланс (ив)"),
    "currency": ("currency", "валюта", "curr"),
    "card_type": ("cardtype", "bp type", "тип бп", "тип"),
}
_LEGAL = re.compile(r"\b(ооо|мчж|mchj|llc|ltd|xk|хк|ип|ак|ao|ао|оао|зао|чп|xususiy|корхонаси|корхона|фирма|firm|"
                    r"компания|company|group|груп|sp|сп|qk|кк|dk|дк|unitar|unitary|davlat)\b")


# ------------------------------------------------------------------ helpers


def digits(value: Any) -> str:
    return re.sub(r"\D", "", str(value or ""))


def norm_name(name: Any) -> str:
    """A counterparty's name for matching: Cyrillic, lower case, no legal form or punctuation."""
    from integrations.common.translit import latin_to_cyrillic

    text = str(name or "").lower().replace("ё", "е")
    text = latin_to_cyrillic(text) if re.search(r"[a-z]", text) else text
    text = re.sub(r"[\"'«»“”‘’`.,()\-–—/\\]+", " ", text)
    text = _LEGAL.sub(" ", text)
    return " ".join(text.split())


def num(value: Any) -> float:
    if value in (None, ""):
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).replace(" ", "").replace(" ", "")
    if text.count(",") == 1 and text.count(".") == 0:
        text = text.replace(",", ".")
    text = text.replace(",", "")
    try:
        return float(text)
    except ValueError:
        return 0.0


# ------------------------------------------------------------------ 1C


@dataclass
class OneCParty:
    key: str
    name: str
    inn: str
    payable: float = 0.0          # 6010 + 6015, so'm (credit balance = we owe)
    payable_fc: dict[str, float] = field(default_factory=dict)  # 6015 in its currency
    advance: float = 0.0          # 4310 + 4315, so'm (debit = we paid ahead)
    lines: list[dict[str, Any]] = field(default_factory=list)

    @property
    def unoffset(self) -> bool:
        """Debt and an advance on the same supplier at once: the advance may never have been offset (зачет)."""
        return self.payable > TOLERANCE_SOM and self.advance > TOLERANCE_SOM

    @property
    def net(self) -> float:
        """What we owe after the advances: positive = we owe the supplier."""
        return self.payable - self.advance


def onec_parties(data: dict[str, Any]) -> dict[str, OneCParty]:
    """1C's register rows → one record per counterparty (only counterparties, not persons)."""
    accounts = {a["Ref_Key"]: (str(a.get("Code", "")).strip(), a.get("Description", "")) for a in data["chart"]}
    parties = {p["Ref_Key"]: p for p in data["parties"]}
    contracts = {c["Ref_Key"]: c.get("Description", "") for c in data.get("contracts", [])}
    currencies = {c["Ref_Key"]: c.get("Description", "") for c in data.get("currencies", [])}
    found: dict[str, OneCParty] = {}
    for row in data["rows"]:
        code, account_name = accounts.get(row.get("Account_Key"), ("", ""))
        if code not in PAYABLE + ADVANCE or not str(row.get("ExtDimension1_Type", "")).endswith("Catalog_Контрагенты"):
            continue
        key = row.get("ExtDimension1") or ""
        info = parties.get(key, {})
        party = found.setdefault(key, OneCParty(key, info.get("Description") or "(номи йўқ)", digits(info.get("ИНН"))))
        cr, dr = num(row.get("СуммаBalanceCr")), num(row.get("СуммаBalanceDr"))
        if code in PAYABLE:
            party.payable += cr - dr
            if code == "6015":
                cur = currencies.get(row.get("Валюта_Key"), "валюта")
                party.payable_fc[cur] = party.payable_fc.get(cur, 0.0) + num(row.get("ВалютнаяСуммаBalanceCr")) - num(row.get("ВалютнаяСуммаBalanceDr"))
        else:
            party.advance += dr - cr
        party.lines.append({"account": code, "account_name": account_name,
                            "contract": contracts.get(row.get("ExtDimension2"), ""), "dr": dr, "cr": cr,
                            "currency": currencies.get(row.get("Валюта_Key"), ""),
                            "fc": num(row.get("ВалютнаяСуммаBalanceCr")) - num(row.get("ВалютнаяСуммаBalanceDr"))})
    return found


async def pull_onec(out: Path) -> dict[str, int]:
    """Read what the comparison needs from 1C (GET only) into ``out``; returns counts, never amounts."""
    from integrations.onec.client import OneCClient

    async with OneCClient(agent="ap-reconcile") as client:
        chart = (await client.get("ChartOfAccounts_Хозрасчетный", {"$select": "Ref_Key,Code,Description"})).get("value", [])
        chart = [a for a in chart if str(a.get("Code", "")).startswith(("60", "43"))]
        keys = {a["Ref_Key"] for a in chart}
        moment = datetime.now().strftime("%Y-%m-%dT%H:%M:%S")
        rows = (await client.get(f"AccountingRegister_Хозрасчетный/Balance(Period=datetime'{moment}')")).get("value", [])
        data = {
            "as_of": moment,
            "chart": chart,
            "rows": [r for r in rows if r.get("Account_Key") in keys],
            "parties": (await client.get("Catalog_Контрагенты", {"$select": "Ref_Key,Description,ИНН"})).get("value", []),
            "contracts": (await client.get("Catalog_ДоговорыКонтрагентов", {"$select": "Ref_Key,Description"})).get("value", []),
            "currencies": (await client.get("Catalog_Валюты", {"$select": "Ref_Key,Code,Description"})).get("value", []),
        }
    out.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
    return {"rows": len(data["rows"]), "parties": len(data["parties"])}


# ------------------------------------------------------------------ SAP


@dataclass
class SapParty:
    code: str
    name: str
    inn: str
    owed: float          # so'm: what we owe the supplier (positive)
    owed_lc: float       # SAP's local currency (USD)
    currency: str = ""
    owed_fc: float = 0.0
    som_known: bool = True


def read_table(path: Path) -> list[dict[str, Any]]:
    """An Excel / CSV export as rows of {header: value}."""
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        sheet = load_workbook(path, read_only=True, data_only=True).active
        rows = [list(r) for r in sheet.iter_rows(values_only=True)]
    else:
        text = path.read_text(encoding="utf-8-sig")
        dialect = csv.Sniffer().sniff(text[:4000], delimiters=",;\t")
        rows = list(csv.reader(text.splitlines(), dialect))
    header_at = next((i for i, r in enumerate(rows[:20]) if sum(1 for c in r if str(c or "").strip()) >= 3), 0)
    header = [str(c or "").strip() for c in rows[header_at]]
    return [dict(zip(header, r)) for r in rows[header_at + 1:] if any(str(c or "").strip() for c in r)]


def find_columns(header: list[str]) -> dict[str, str]:
    """Our field → the export's column, by header name (exact first, then contains)."""
    lowered = {h.lower().strip(): h for h in header if h}
    found: dict[str, str] = {}
    for key, names in SAP_HEADERS.items():
        exact = next((lowered[n] for n in names if n in lowered and lowered[n] not in found.values()), None)
        if exact is None:
            exact = next((h for low, h in lowered.items() for n in names
                          if n in low and h not in found.values() and len(n) > 3), None)
        if exact:
            found[key] = exact
    return found


def sap_parties(rows: list[dict[str, Any]]) -> tuple[dict[str, SapParty], dict[str, Any]]:
    """The export → suppliers, with what we owe each in so'm when SAP gives it (its system currency)."""
    if not rows:
        return {}, {"columns": {}, "flipped": False}
    columns = find_columns(list(rows[0]))
    if "name" not in columns or not ({"balance_sys", "balance_lc"} & set(columns)):
        raise ValueError(f"SAP export: need a name and a balance column; found {sorted(columns)}")
    kept = [r for r in rows if not ("card_type" in columns and str(r.get(columns["card_type"], "")).strip().upper()
                                    not in ("S", "SUPPLIER", "ПОСТАВЩИК", "ПОСТ."))]
    values = [num(r.get(columns.get("balance_sys") or columns["balance_lc"])) for r in kept]
    # SAP stores a supplier's balance as a credit: negative = we owe. Shown positive here.
    flipped = sum(1 for v in values if v < 0) > sum(1 for v in values if v > 0)
    sign = -1.0 if flipped else 1.0
    found: dict[str, SapParty] = {}
    for r in kept:
        name = str(r.get(columns["name"], "") or "").strip()
        if not name:
            continue
        code = str(r.get(columns.get("code", ""), "") or "").strip() or name
        som_known = "balance_sys" in columns
        lc = sign * num(r.get(columns.get("balance_lc", ""), 0))
        owed = sign * num(r.get(columns["balance_sys"])) if som_known else lc
        found[code] = SapParty(code, name, digits(r.get(columns.get("inn", ""), "")), owed, lc,
                               str(r.get(columns.get("currency", ""), "") or "").strip(),
                               sign * num(r.get(columns.get("balance_fc", ""), 0)), som_known)
    return found, {"columns": columns, "flipped": flipped}


# ------------------------------------------------------------------ compare


@dataclass
class Match:
    onec: OneCParty | None
    sap: SapParty | None
    by: str = ""

    @property
    def diff(self) -> float:
        return (self.onec.net if self.onec else 0.0) - (self.sap.owed if self.sap else 0.0)


def tolerance(amount: float) -> float:
    return max(TOLERANCE_SOM, abs(amount) * TOLERANCE_SHARE)


def match(onec: dict[str, OneCParty], sap: dict[str, SapParty]) -> list[Match]:
    """Pair counterparties: by ИНН first, then by the cleaned-up name; the rest stay alone."""
    left = dict(onec)
    pairs: list[Match] = []
    by_inn = {p.inn: k for k, p in left.items() if len(p.inn) >= 9}
    by_name: dict[str, str] = {}
    for k, p in left.items():
        by_name.setdefault(norm_name(p.name), k)
    for s in sorted(sap.values(), key=lambda s: -abs(s.owed)):
        key, how = None, ""
        if len(s.inn) >= 9 and by_inn.get(s.inn) in left:
            key, how = by_inn[s.inn], "ИНН"
        elif by_name.get(norm_name(s.name)) in left:
            key, how = by_name[norm_name(s.name)], "ном"
        pairs.append(Match(left.pop(key) if key else None, s, how))
    pairs += [Match(p, None) for p in left.values()]
    return pairs


def reason(m: Match) -> str:
    """The likely reason for a difference, in Uzbek Cyrillic — a hint for the accountant, not a verdict."""
    o, s = m.onec, m.sap
    if o is None:
        return "Фақат SAP'да: 1C'да бу контрагент бўйича қолдиқ йўқ — ҳужжат 1C'га киритилмаган ёки ёпилган"
    if s is None:
        return "Фақат 1C'да: SAP'да бу етказиб берувчи бўйича қолдиқ йўқ — ҳужжат SAP'га киритилмаган ёки контрагент бошқа номда"
    d = m.diff
    if abs(d) <= tolerance(max(abs(o.net), abs(s.owed))):
        return "Мос"
    if o.advance and abs(abs(d) - o.advance) <= tolerance(o.advance):
        return "Аванс фарқи: 1C'да тўланган аванс (4310) ҳисобга олинган, SAP'да эмас (ёки аксинча)"
    if o.unoffset and abs(abs(d) - o.advance) > tolerance(o.advance):
        return UNOFFSET + " — SAP битта соф қолдиқ кўрсатади"
    if (o.net > 0) != (s.owed > 0) and o.net and s.owed:
        return "Йўналиш ҳар хил: бир тизимда қарзмиз, бошқасида ортиқча тўлов/аванс — тўлов бир тизимда киритилмаган"
    if not s.som_known:
        return "SAP суммаси сўмда эмас (USD): курс фарқи бўлиши мумкин — SAP'дан сўмдаги қолдиқ (BalanceSys) керак"
    if o.payable_fc and abs(d) <= 0.05 * max(abs(o.net), abs(s.owed)):
        return "Валюта курси: валютадаги қарз (6015) икки тизимда турли курсда қайта баҳоланган"
    return "Сумма фарқли: бир тизимда ҳужжат ёки тўлов йўқ ёки суммаси бошқа — акт-сверка қилиш керак"


# ------------------------------------------------------------------ workbook


def write_workbook(pairs: list[Match], onec: dict[str, OneCParty], meta: dict[str, Any], out: Path) -> dict[str, int]:
    """The Excel file for the Director and accounting; returns counts (never amounts)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="E8EEF4")
    bad_fill = PatternFill("solid", fgColor="FDECEC")
    money = '#,##0;[Red]-#,##0'

    def sheet(title: str, header: list[str], widths: list[int]):
        ws = wb.create_sheet(title)
        ws.append(header)
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = w
            ws.cell(1, i).font, ws.cell(1, i).fill = bold, head_fill
            ws.cell(1, i).alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A2"
        return ws

    wb.remove(wb.active)
    both = [m for m in pairs if m.onec and m.sap]
    ws = sheet("Солиштириш", ["Контрагент (1C)", "Контрагент (SAP)", "ИНН", "Боғланди", "1C: қарз (6010+6015), сўм",
                              "1C: аванс (4310+4315), сўм", "1C: нетто, сўм", "SAP: қарз, сўм", "SAP: қарз, USD",
                              "Фарқ (1C − SAP), сўм", "Эҳтимолий сабаб"], [34, 34, 13, 10, 18, 18, 16, 16, 14, 16, 60])
    for m in sorted(both, key=lambda m: -abs(m.diff)):
        why = reason(m)
        ws.append([m.onec.name, m.sap.name, m.onec.inn or m.sap.inn, m.by, m.onec.payable, m.onec.advance, m.onec.net,
                   m.sap.owed, m.sap.owed_lc, m.diff, why])
        if why != "Мос":
            for cell in ws[ws.max_row]:
                cell.fill = bad_fill
    for col in "EFGHIJ":
        for cell in ws[col][1:]:
            cell.number_format = money

    only_1c = [m.onec for m in pairs if m.onec and not m.sap and abs(m.onec.net) > TOLERANCE_SOM]
    ws = sheet("Фақат 1C", ["Контрагент", "ИНН", "Қарз (6010+6015), сўм", "Аванс (4310+4315), сўм", "Нетто, сўм",
                            "Валютада (6015)", "Эҳтимолий сабаб", "1C ичида"], [40, 13, 18, 18, 16, 22, 60, 44])
    for p in sorted(only_1c, key=lambda p: -abs(p.net)):
        ws.append([p.name, p.inn, p.payable, p.advance, p.net,
                   "; ".join(f"{v:,.2f} {c}" for c, v in p.payable_fc.items()),
                   reason(Match(p, None)) if meta.get("sap_rows") else "SAP маълумоти ҳали келмаган — солиштирилмади",
                   UNOFFSET if p.unoffset else ""])
    for col in "CDE":
        for cell in ws[col][1:]:
            cell.number_format = money

    only_sap = [m.sap for m in pairs if m.sap and not m.onec and abs(m.sap.owed) > TOLERANCE_SOM]
    ws = sheet("Фақат SAP", ["Код", "Контрагент", "ИНН", "Қарз, сўм", "Қарз, USD", "Валюта", "Эҳтимолий сабаб"],
               [14, 40, 13, 16, 14, 10, 60])
    for s in sorted(only_sap, key=lambda s: -abs(s.owed)):
        ws.append([s.code, s.name, s.inn, s.owed, s.owed_lc, s.currency, reason(Match(None, s))])
    for col in "DE":
        for cell in ws[col][1:]:
            cell.number_format = money

    ws = sheet("1C тафсилот", ["Контрагент", "Ҳисобварақ", "Номи", "Шартнома", "Дт қолдиқ, сўм", "Кт қолдиқ, сўм",
                               "Валютада", "Валюта"], [36, 10, 36, 36, 16, 16, 14, 8])
    for p in sorted(onec.values(), key=lambda p: p.name):
        for line in p.lines:
            ws.append([p.name, line["account"], line["account_name"], line["contract"], line["dr"], line["cr"],
                       line["fc"] or None, line["currency"] if line["fc"] else ""])
    for col in "EF":
        for cell in ws[col][1:]:
            cell.number_format = money

    ws = sheet("Изоҳ", ["Изоҳ"], [120])
    for line in (
        f"1C қолдиқлари: {meta.get('onec_as_of', '')} ҳолатига (OData, фақат ўқиш). SAP: {meta.get('sap_file', '')}.",
        "1C: қарз = 6010 + 6015 ҳисобварақларининг кредит қолдиғи; аванс = 4310 + 4315 дебет қолдиғи; нетто = қарз − аванс.",
        "SAP: етказиб берувчилар қолдиғи (OCRD, тури S); сўмда — SAP'нинг тизим валютаси (BalanceSys), бўлмаса USD.",
        "Боғлаш: аввал ИНН бўйича, кейин тозаланган ном бўйича (ООО, МЧЖ, қўштирноқлар олиб ташланган).",
        f"Фарқ кичик бўлса (≤ {TOLERANCE_SOM:,.0f} сўм ёки 0,5 %) «Мос» деб олинади.",
        "Сабаб — тахмин, ҳукм эмас: ҳар бир фарқ бухгалтерия томонидан акт-сверка билан текширилади.",
        f"SAP устунлари топилди: {', '.join(f'{k}={v}' for k, v in meta.get('sap_columns', {}).items())}"
        + ("; SAP қолдиқлари манфий эди — ишораси алмаштирилди" if meta.get("flipped") else ""),
    ):
        ws.append([line])
    wb.save(out)
    return {"matched": len(both), "differ": sum(1 for m in both if reason(m) != "Мос"),
            "only_1c": len(only_1c), "only_sap": len(only_sap)}


def main() -> None:
    parser = argparse.ArgumentParser(description="1C vs SAP supplier debt, read-only")
    parser.add_argument("--pull-1c", type=Path, help="read 1C now (GET only) into this JSON file")
    parser.add_argument("--onec-json", type=Path, help="1C data read earlier (--pull-1c)")
    parser.add_argument("--sap", type=Path, help="SAP B1 export of the suppliers' balances (xlsx or csv)")
    parser.add_argument("--out", type=Path, default=Path("kreditorlik-1c-sap.xlsx"))
    args = parser.parse_args()
    if args.pull_1c:
        print("1C read:", asyncio.run(pull_onec(args.pull_1c)))
        return
    data = json.loads(args.onec_json.read_text(encoding="utf-8"))
    onec = onec_parties(data)
    sap, info = sap_parties(read_table(args.sap)) if args.sap else ({}, {"columns": {}, "flipped": False})
    counts = write_workbook(match(onec, sap), onec, {"onec_as_of": data.get("as_of", ""),
                                                     "sap_file": args.sap.name if args.sap else "йўқ",
                                                     "sap_columns": info["columns"], "flipped": info["flipped"],
                                                     "sap_rows": len(sap)},
                            args.out)
    print("written", args.out, counts)  # counts only, never amounts


if __name__ == "__main__":
    main()
