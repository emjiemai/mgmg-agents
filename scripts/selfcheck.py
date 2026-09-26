"""Offline self-check for the pure logic — no network, no database.

Verifies the parts that would otherwise only be exercised against live SAP,
the CRM and Telegram: money formatting, aging buckets, Telegram message
splitting, org_bot parsing, and brief/receivables rendering.

Run:
    python scripts/selfcheck.py
"""

from __future__ import annotations

import sys
from datetime import date, datetime, timedelta
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from integrations.common.money import (
    format_money,
    format_money_by_currency,
    format_uzs,
    format_uzs_short,
    from_tiyin,
    to_tiyin,
    uzs_to_usd,
)
from integrations.common.timeutil import TASHKENT, days_between, parse_sap_date, to_local, to_utc
from integrations.org_bot.ops_manager import (
    _task_card_text,
    _task_keyboard,
    parse_callback_data,
    parse_role_and_request,
    validate_classification,
)
from integrations.org_bot.roles import AGENT_SLUGS, ROLE_SLUGS
from integrations.telegram.bot import sanitize_model_html
from integrations.sap.aging import aging_bucket
from integrations.telegram.bot import escape, split_message

FAILURES: list[str] = []


def check(label: str, actual: object, expected: object) -> None:
    """Assert equality and record the outcome.

    Args:
        label: Name of the case.
        actual: Value produced.
        expected: Value required.
    """
    if actual == expected:
        print(f"  ok   {label}")
    else:
        print(f"  FAIL {label}: got {actual!r}, want {expected!r}")
        FAILURES.append(label)


def check_true(label: str, condition: bool) -> None:
    """Assert a condition holds.

    Args:
        label: Name of the case.
        condition: Must be True.
    """
    check(label, bool(condition), True)


def latin_words(text: str, allow: set[str] | None = None) -> list[str]:
    """Latin-script words left in a message that should be Uzbek Cyrillic.

    Tags, brand names and codes are fine; anything else is a missed
    translation (the business's rule since 2026-09-25).
    """
    import re

    allowed = {"CEO", "IT", "KPI", "SAP", "CRM", "HR", "AI", "Garmin", "EMJ", "SOP", "ADM", "OPS", "Admin", "Bot"}
    allowed |= allow or set()
    # Telegram commands (/ismlar, /bekor) can only be Latin.
    plain = re.sub(r"/[a-z_]+", " ", re.sub(r"<[^>]+>", " ", text))
    return [w for w in re.findall(r"[A-Za-z][A-Za-z0-9]*(?:'[A-Za-z]+)*", plain) if w not in allowed]


def test_references() -> None:
    """Every settings.X and every attribute of a project module that the code uses exists.

    2026-09-26: a removed setting (ops_manager_bot_provider) was still read
    in the answer path, so every question to OPS Manager Bot failed with
    "Хатолик юз берди" — invisible to lint and imports, because Python only
    looks an attribute up when that line runs. This walks every file instead.
    Scope-aware: an import inside one function applies only there, and a
    local variable with the same name as a module hides it.
    """
    print("references")
    import ast
    import importlib

    from integrations.common.config import PROJECT_ROOT, settings

    scope_types = (ast.FunctionDef, ast.AsyncFunctionDef, ast.Lambda, ast.ClassDef)

    def own_nodes(scope: ast.AST):
        """Nodes of this scope, not of the functions nested inside it."""
        stack = list(ast.iter_child_nodes(scope))
        while stack:
            node = stack.pop()
            yield node
            if not isinstance(node, scope_types):
                stack.extend(ast.iter_child_nodes(node))

    def project_modules(nodes) -> dict[str, object]:
        found: dict[str, object] = {}
        for node in nodes:
            if isinstance(node, ast.ImportFrom) and node.module and node.module.startswith("integrations"):
                for alias in node.names:
                    try:  # only names that are themselves modules ("from x import store")
                        found[alias.asname or alias.name] = importlib.import_module(f"{node.module}.{alias.name}")
                    except ImportError:
                        pass
            elif isinstance(node, ast.Import):
                for alias in node.names:
                    if alias.name.startswith("integrations") and alias.asname:
                        found[alias.asname] = importlib.import_module(alias.name)
        return found

    missing: list[str] = []
    files = [
        p for folder in ("integrations", "agents", "scripts")
        for p in (PROJECT_ROOT / folder).rglob("*.py") if "__pycache__" not in p.parts
    ]
    for path in files:
        tree = ast.parse(path.read_text(encoding="utf-8"))
        module_level = project_modules(own_nodes(tree))
        scopes = [tree] + [n for n in ast.walk(tree) if isinstance(n, (ast.FunctionDef, ast.AsyncFunctionDef))]
        for scope in scopes:
            nodes = list(own_nodes(scope))
            local_imports = project_modules(nodes) if scope is not tree else {}
            stored = {n.id for n in nodes if isinstance(n, ast.Name) and isinstance(n.ctx, ast.Store)}
            if scope is not tree:
                stored |= {a.arg for a in scope.args.args + scope.args.kwonlyargs}
            modules = {**{k: v for k, v in module_level.items() if k not in stored}, **local_imports}
            for node in nodes:
                if not (isinstance(node, ast.Attribute) and isinstance(node.value, ast.Name)):
                    continue
                owner, attr = node.value.id, node.attr
                where = f"{path.relative_to(PROJECT_ROOT)}:{node.lineno}"
                if owner == "settings" and "settings" not in stored and not hasattr(settings, attr):
                    missing.append(f"{where} settings.{attr}")
                elif owner in modules and not hasattr(modules[owner], attr):
                    missing.append(f"{where} {owner}.{attr}")
    check("no reference to a setting or function that doesn't exist", missing, [])


def test_bot_flows() -> None:
    """The Director's real message paths, end to end, with AI/Telegram/DB faked.

    Added 2026-09-26 after every question failed with "Хатолик юз берди": the
    answer path had never been executed by any test. This runs it — and the
    task path, and every data agent's fetcher on an empty database — so a
    runtime error anywhere in them fails here instead of in production.
    """
    print("bot flows (smoke)")
    import asyncio
    import sys
    import uuid

    from integrations.common import db
    from integrations.common.agent_loader import load_agent
    from integrations.org_bot import ops_manager, roles, store

    load_agent("cash-calendar")  # loaded now so its database functions get faked below

    async def no_rows(*_args, **_kwargs):
        return []

    async def no_row(*_args, **_kwargs):
        return None

    director = {"id": "d", "telegram_user_id": 1, "role": "operatsion_direktor", "status": "active",
                "display_name": "Director", "full_name": "Бобур Алиев"}
    worker = {"id": "w", "telegram_user_id": 2, "role": "it", "status": "active",
              "display_name": "GMHRD", "full_name": "Алишер Каримов"}

    async def employees_by_role(role):
        return [director] if role == "operatsion_direktor" else [worker] if role == "it" else []

    async def active_employees():
        return [director, worker]

    async def new_task(**kwargs):
        return {"id": "t1", **kwargs}

    replies: list[str] = []

    async def fake_reply(_chat_id, _run_id, text, reply_markup=None):
        replies.append(text)
        return [100]

    class FakeAI:
        classification: dict = {}

        def __init__(self, *args, **kwargs):  # accepts exactly what the real client accepts
            import inspect

            from integrations.ai.openrouter_client import OpenRouterClient

            inspect.signature(OpenRouterClient.__init__).bind(self, *args, **kwargs)

        async def __aenter__(self):
            return self

        async def __aexit__(self, *exc):
            return None

        async def complete_json(self, system, user):
            return dict(FakeAI.classification)

        async def complete(self, system, user, **kwargs):
            return "Жавоб"

    class FakeBot:
        def __init__(self, *args, **kwargs):
            pass

        async def __aenter__(self):
            return self

        async def __aexit__(self, *exc):
            return None

        async def send_message(self, text, **kwargs):
            replies.append(text)
            return [200]

        async def _call(self, *args, **kwargs):
            return {"message_id": 201}

    class NoSheets:
        def __init__(self, *args, **kwargs):
            pass

        async def __aenter__(self):
            from integrations.google.sheets_client import SheetsError

            raise SheetsError("offline")

        async def __aexit__(self, *exc):
            return None

    # Fake every database helper wherever it was imported by name.
    saved: list[tuple[object, str, object]] = []

    def patch(obj, name, value):
        saved.append((obj, name, getattr(obj, name)))
        setattr(obj, name, value)

    fakes = {"fetch_all": no_rows, "fetch_one": no_row, "execute": no_row}
    # Captured first: patching db itself mid-loop must not change what we match.
    originals = {name: getattr(db, name) for name in fakes}
    for module in list(sys.modules.values()):
        module_name = getattr(module, "__name__", "")
        if not (module_name.startswith("integrations") or module_name.startswith("mgmg_agent_")):
            continue
        for name, fake in fakes.items():
            if getattr(module, name, None) is originals[name]:
                patch(module, name, fake)
    for name in dir(store):
        if name.startswith("_") or not asyncio.iscoroutinefunction(getattr(store, name)):
            continue
        patch(store, name, no_row if name.startswith(("get_", "create_", "find_", "set_", "mark_", "log_")) else no_rows)
    patch(store, "active_employees_by_role", employees_by_role)
    patch(store, "list_active_employees", active_employees)
    patch(store, "create_task", new_task)
    patch(ops_manager, "OpenRouterClient", FakeAI)
    patch(ops_manager, "TelegramBot", FakeBot)
    patch(ops_manager, "SheetsClient", NoSheets)
    patch(ops_manager, "_reply", fake_reply)

    try:
        for slug in sorted(roles.AGENT_SLUGS):
            try:
                data = asyncio.run(ops_manager._fetch_agent_data(slug))
                ok = isinstance(data, str) and bool(data)
            except Exception as exc:  # noqa: BLE001 — the point is to report it
                ok = False
                print(f"    {slug}: {type(exc).__name__}: {exc}")
            check_true(f"agent '{slug}' answers from an empty database", ok)

        FakeAI.classification = {"target_type": "agent", "target_agent": "xodimlar_kpi", "task_summary": "",
                                 "target_employee": None, "due_date": None}
        replies.clear()
        asyncio.run(ops_manager._dispatch_director_task(1, "kechagi hisobotlarni korsat", 11, uuid.uuid4()))
        check("a question gets the AI's answer, not the error", replies, ["Жавоб"])

        FakeAI.classification = {"target_type": "employee", "target_role": "it", "target_employee": "E1",
                                 "task_summary": "Принтерни текширинг", "due_date": None}
        replies.clear()
        asyncio.run(ops_manager._dispatch_director_task(1, "Alisherga ayt printerni tekshirsin", 12, uuid.uuid4()))
        check_true("a task reaches the named person", any("Принтерни текширинг" in r for r in replies))
        check_true("and the Director is told who got it", any(r.startswith("Юборилди: Алишер Каримов") for r in replies))
        check_true("no error on the task path", not any("Хатолик" in r for r in replies))
    finally:
        for obj, name, value in reversed(saved):
            setattr(obj, name, value)


def test_money() -> None:
    """Money conversion and Uzbek sum formatting."""
    print("money")
    nbsp = " "
    check("to_tiyin whole", to_tiyin(1250), 125000)
    check("to_tiyin decimal", to_tiyin("1250.75"), 125075)
    check("to_tiyin rounds half up", to_tiyin(0.005), 1)
    check("to_tiyin None", to_tiyin(None), 0)
    check("to_tiyin garbage", to_tiyin("n/a"), 0)
    check("from_tiyin", str(from_tiyin(125075)), "1250.75")
    check("format thousands", format_uzs(125000000), f"1{nbsp}250{nbsp}000{nbsp}сўм")
    check("format negative", format_uzs(-125000000), f"-1{nbsp}250{nbsp}000{nbsp}сўм")
    check("format no currency", format_uzs(125000000, with_currency=False), f"1{nbsp}250{nbsp}000")
    check("short mlrd", format_uzs_short(int(1.25e9 * 100)), f"1,25{nbsp}млрд{nbsp}сўм")
    check("short mln", format_uzs_short(340_000_000 * 100), f"340{nbsp}млн{nbsp}сўм")
    check("short small", format_uzs_short(85_000 * 100), f"85{nbsp}000{nbsp}сўм")
    check("short negative", format_uzs_short(-int(2e9 * 100)), f"-2{nbsp}млрд{nbsp}сўм")
    check("usd reference", str(uzs_to_usd(12_800 * 100)), "1.00")

    # 2026-09-07: SAP AR invoices turned out to include real USD-denominated
    # ones that format_uzs/format_uzs_short were silently mislabeling as
    # so'm (right number, wrong currency). format_money/format_money_by_currency
    # replace every such call site in receivables + ceo-daily-brief.
    check("format_money UZS unchanged", format_money(125000000, "UZS"), format_uzs(125000000))
    check("format_money USD", format_money(976400, "USD"), "$9,764.00")
    check("format_money None currency defaults UZS", format_money(125000000, None), format_uzs(125000000))
    check("format_money unknown currency code", format_money(150000, "EUR"), "1,500.00 EUR")
    check(
        "format_money_by_currency single currency == format_money",
        format_money_by_currency([(125000000, "UZS")]),
        format_uzs(125000000),
    )
    check(
        "format_money_by_currency sums within a currency, not across",
        format_money_by_currency([(1_000_000, "UZS"), (500_000, "UZS"), (976_400, "USD")]),
        f"15{nbsp}000{nbsp}сўм + $9,764.00",
    )
    check("format_money_by_currency empty", format_money_by_currency([]), format_uzs(0))


def test_time() -> None:
    """Timezone conversion and SAP date parsing."""
    print("time")
    utc_noon = datetime(2026, 8, 18, 7, 0, tzinfo=to_utc(datetime(2026, 1, 1)).tzinfo)
    check("utc -> tashkent is +5", to_local(utc_noon).hour, 12)
    naive_local = datetime(2026, 8, 18, 9, 0)
    check("naive treated as tashkent", to_utc(naive_local).hour, 4)
    check("tashkent offset", TASHKENT.utcoffset(datetime(2026, 8, 18)), timedelta(hours=5))
    check("parse date", parse_sap_date("2026-08-18"), date(2026, 8, 18))
    check("parse datetime string", parse_sap_date("2026-08-18T00:00:00Z"), date(2026, 8, 18))
    check("parse empty", parse_sap_date(""), None)
    check("parse garbage", parse_sap_date("not-a-date"), None)
    check("days_between", days_between(date(2026, 8, 1), date(2026, 8, 18)), 17)
    check("days_between None", days_between(None), 0)


def test_aging() -> None:
    """AR aging bucket boundaries."""
    print("aging buckets")
    check("not due", aging_bucket(0), "current")
    check("future", aging_bucket(-5), "current")
    check("1 day", aging_bucket(1), "1_30")
    check("30 days", aging_bucket(30), "1_30")
    check("31 days", aging_bucket(31), "31_60")
    check("60 days", aging_bucket(60), "31_60")
    check("61 days", aging_bucket(61), "61_90")
    check("90 days", aging_bucket(90), "61_90")
    check("91 days", aging_bucket(91), "90_plus")


def test_telegram() -> None:
    """Message splitting and HTML escaping."""
    print("telegram")
    check("short message not split", len(split_message("hello")), 1)

    long_text = "\n".join(f"line {i} " + "x" * 100 for i in range(200))
    chunks = split_message(long_text)
    check_true("long message split", len(chunks) > 1)
    check_true("every chunk within limit", all(len(c) <= 3900 for c in chunks))
    check("no content lost", sum(c.count("line ") for c in chunks), 200)

    single_line = "y" * 9000
    hard = split_message(single_line)
    check_true("pathological line hard-cut", all(len(c) <= 3900 for c in hard))
    check("hard-cut keeps all characters", sum(len(c) for c in hard), 9000)

    check("escape ampersand", escape("Rogers & Co <Ltd>"), "Rogers &amp; Co &lt;Ltd&gt;")
    check("escape None", escape(None), "")


def test_org_bot() -> None:
    """OPS Manager Bot's callback parsers and classification validator."""
    print("org_bot")

    check("callback: setrole splits once", parse_callback_data("setrole:it:req-1"), ("setrole", "it:req-1"))
    check("callback: taskdone splits once", parse_callback_data("taskdone:task-1"), ("taskdone", "task-1"))
    check("callback: no colon is malformed", parse_callback_data("garbage"), None)

    check("role+request: splits once", parse_role_and_request("it:req-1"), ("it", "req-1"))
    check("role+request: id itself may contain colons", parse_role_and_request("it:req:1"), ("it", "req:1"))
    check("role+request: no colon is malformed", parse_role_and_request("it"), None)

    check_true("every role slug is a known role", all(r in ROLE_SLUGS for r in ("it", "hr", "ombor")))
    check_true("bogus role slug is rejected", "not_a_role" not in ROLE_SLUGS)

    # 2026-09-16: a picked role now needs the admin's second Accept.
    from integrations.org_bot.admin import role_decision_keyboard

    role_buttons = role_decision_keyboard("3fa85f64-5717-4562-b3fc-2c963f66afa6")["inline_keyboard"][0]
    check(
        "role request card has Accept + Reject",
        [b["callback_data"].split(":", 1)[0] for b in role_buttons],
        ["role_approve", "role_reject"],
    )
    check_true(
        "role request callback fits Telegram's 64-byte limit",
        all(len(b["callback_data"].encode()) <= 64 for b in role_buttons),
    )

    check(
        "classify: valid employee route",
        validate_classification({"target_type": "employee", "target_role": "it", "target_agent": None}),
        ("employee", "it", None),
    )
    check(
        "classify: valid agent route",
        validate_classification({"target_type": "agent", "target_role": None, "target_agent": "lead_agent"}),
        ("agent", None, "lead_agent"),
    )
    check(
        "classify: explicit none is valid, not an error",
        validate_classification({"target_type": "none"}),
        ("none", None, None),
    )
    check(
        "classify: refused (guardrail path) is valid, not an error",
        validate_classification({"target_type": "refused", "task_summary": "Kechirasiz, bunga yordam bera olmayman."}),
        ("refused", None, None),
    )
    check(
        "classify: employee with unknown role slug is rejected",
        validate_classification({"target_type": "employee", "target_role": "made_up_role"}),
        None,
    )
    check(
        "classify: agent with unknown agent slug is rejected",
        validate_classification({"target_type": "agent", "target_agent": "made_up_agent"}),
        None,
    )
    check(
        "classify: employee route missing its role is rejected",
        validate_classification({"target_type": "employee", "target_role": None}),
        None,
    )
    check("classify: malformed response is rejected", validate_classification({}), None)
    check_true("agent vocabulary is non-empty", len(AGENT_SLUGS) > 0)

    check_true("task card text carries the summary", "Ombor" in _task_card_text("Ombor tekshiruvi"))
    fresh_kb = _task_keyboard("task-1")
    check("fresh task has Start + Done buttons", len(fresh_kb["inline_keyboard"][0]), 2)
    started_kb = _task_keyboard("task-1", started=True)
    check("started task has only the Done button", len(started_kb["inline_keyboard"][0]), 1)
    check_true("Start button carries the right callback_data", "taskstart:task-1" in str(fresh_kb))
    check_true("Done button carries the right callback_data", "taskdone:task-1" in str(fresh_kb))
    check_true("started keyboard has no Start button left", "taskstart:" not in str(started_kb))

    # The real bug: escape() alone turns a model's deliberate <b> into
    # literal visible "<b>" text once Telegram's HTML parser unescapes
    # &lt;b&gt; back to a literal "<" for display. sanitize_model_html must
    # let the allowed tags through as real tags while still neutralizing
    # anything else, so a stray "<" from underlying data can't break the
    # HTML parse or render as unintended markup.
    check(
        "sanitize: allowed <b> survives as a real tag",
        sanitize_model_html("Eng yaxshi lid: <b>JW Marriott</b>"),
        "Eng yaxshi lid: <b>JW Marriott</b>",
    )
    check(
        "sanitize: allowed <i> survives as a real tag",
        sanitize_model_html("<i>Eslatma</i>: tekshiring"),
        "<i>Eslatma</i>: tekshiring",
    )
    check(
        "sanitize: an unrelated tag is neutralized, not rendered",
        sanitize_model_html("narx < 100 va <script>alert(1)</script>"),
        "narx &lt; 100 va &lt;script&gt;alert(1)&lt;/script&gt;",
    )
    check("sanitize: None becomes empty string", sanitize_model_html(None), "")


def test_brief_rendering() -> None:
    """The CEO brief: the five numbers (A2), honest gaps, and who didn't report."""
    print("brief rendering (A2)")
    from integrations.common.agent_loader import load_agent
    from integrations.sap.figures import Figure
    from integrations.sap.models import ARAging, ARInvoice

    brief = load_agent("ceo-daily-brief")

    empty = brief.BriefData()
    empty.note_failure("sap_orders", RuntimeError("connection refused"))
    text = brief.render(empty)
    check_true("title and time on one line", "CEO кунлик ҳисоботи" in text.splitlines()[0] and "Тошкент" in text.splitlines()[0])
    check_true("cash says plainly it isn't connected", "💰 Касса: <i>уланмаган</i>" in text)
    check_true("a missing feed reads 'маълумот йўқ', never a number", "📈 Кечаги сотув: <i>маълумот йўқ</i>" in text)
    check_true("all five lines are there", all(k in text for k in ("Касса", "Кечаги сотув", "Захира", "Мижоз қарзи", "Бугунги тўловлар")))

    aging = ARAging(snapshot_date=date(2026, 9, 25))
    aging.invoices = [
        ARInvoice(doc_entry=1, doc_num=1001, card_code="C1", card_name="A", days_overdue=120, aging_bucket="90_plus",
                  currency="USD", doc_total_tiyin=500_000, balance_due_tiyin=500_000),
        ARInvoice(doc_entry=2, doc_num=1002, card_code="C2", card_name="B", days_overdue=0, aging_bucket="current",
                  currency="USD", doc_total_tiyin=1_000_000, balance_due_tiyin=1_000_000),
    ]
    aging.total_overdue_tiyin = 500_000
    aging.bucket_totals_tiyin = {"90_plus": 500_000}
    aging.bucket_counts = {"90_plus": 1}
    data = brief.BriefData(
        aging=aging,
        sales=Figure(status="ok", totals={"USD": 1_234_000}, count=3, as_of=date(2026, 9, 26)),
        inventory=Figure(status="ok", totals={"USD": 12_000_000}, count=100, capped=True, as_of=date(2026, 9, 26)),
        payments=brief.PaymentsDue(totals={"UZS": 1_500_000_000}, count=2, unclear=1),
        report_rows=[],
        previous={"debt": {"status": "ok", "totals": {"USD": 1_300_000}, "capped": False},
                  "sales": {"status": "ok", "totals": {"USD": 1_234_000}, "capped": False}},
    )
    text = brief.render(data)
    check_true("yesterday's sales with the order count", "📈 Кечаги сотув: $12,340.00 (3 та буюртма)" in text)
    check_true("unchanged sales show no change", "Кечаги сотув: $12,340.00 (3 та буюртма)\n" in text)
    check_true("a capped feed is a lower bound", "📦 Захира: камида $120,000.00*" in text)
    check_true("the lower bound is explained", "рақам тўлиқ эмас" in text)
    check_true("debt total, overdue part, 90+ marker", "$15,000.00" in text and "муддати ўтгани $5,000.00 (1 та) 🔴" in text)
    check_true("debt change since the last brief", "кечагига ▲ $2,000.00" in text)
    check_true("today's approved payments", "💳 Бугунги тўловлар: 2 та — 15" in text and "сўм" in text)
    check_true("an unreadable payment date is said, not guessed", "1 та тасдиқланган сўровда сана аниқ эмас" in text)
    check_true("the brief is Uzbek Cyrillic", latin_words(text, allow={"CEO", "SAP"}) == [])
    stored = brief.five_numbers_json(data)
    check("stored for tomorrow's change", stored["debt"]["totals"], {"USD": 1_500_000})
    check("capped flag kept", stored["inventory"]["capped"], True)

    stale = Figure(status="stale", as_of=date(2026, 9, 1))
    check_true("an old feed says since when", "01.09.2026 дан бери янгиланмаган" in brief.render(brief.BriefData(sales=stale)))

    # 2026-09-16: daily reports never reach the Director one by one; the
    # brief names only who didn't report on the last day they were asked.
    asked_day = date(2026, 9, 15)
    rows = [
        {"report_date": asked_day, "status": "submitted", "display_name": "Aziz", "role": "it"},
        {"report_date": asked_day, "status": "asked", "display_name": "Dmitriy", "role": "b2b_sotuv"},
    ]
    text = brief.render(brief.BriefData(report_rows=rows))
    check_true("non-reporter named", "Dmitriy" in text)
    check_true("someone who reported is not named", "Aziz" not in text)
    check_true("missed out of asked is shown", "1 / 2" in text)
    everyone_in = [dict(r, status="submitted") for r in rows]
    check_true(
        "everyone reporting is stated",
        "ҳаммаси юборди (2/2)" in brief.render(brief.BriefData(report_rows=everyone_in)),
    )
    never_asked = brief.render(brief.BriefData(report_rows=[]))
    check_true(
        "nobody asked is said plainly, not 'everyone reported'",
        "сўралмаган" in never_asked and "ҳаммаси юборди" not in never_asked,
    )
    check_true("removed sections stay removed", all(w not in never_asked for w in ("Reportlar", "Pipeline", "Davomat")))

    import asyncio

    from integrations.common.config import settings

    if not settings.daily_reports_enabled:
        check("switched off: no report rows fetched", asyncio.run(brief._fetch_report_results()), [])


def test_dates_and_figures() -> None:
    """Strict date reading, SAP dashboard figures, and payment due days."""
    print("dates and SAP figures")
    from integrations.common.dates import parse_day
    from integrations.org_bot.permissions import due_day
    from integrations.sap import figures

    ref = date(2026, 9, 23)
    for text, want in (
        ("25.09.2026", date(2026, 9, 25)), ("25.09", date(2026, 9, 25)), ("2026-09-25", date(2026, 9, 25)),
        ("25 сентябр", date(2026, 9, 25)), ("25-sentabr", date(2026, 9, 25)), ("эртага", date(2026, 9, 24)),
        ("бугун 15.00", ref), ("3 кун ичида", date(2026, 9, 26)), ("05.01", date(2027, 1, 5)),
        ("шу ҳафта ичида", None), ("тезроқ", None), ("30.02.2026", None), ("5 млн сўм", None),
    ):
        check(f"date '{text}'", parse_day(text, ref), want)

    from datetime import datetime, timezone

    request = {"execute_by": "эртага", "submitted_at": datetime(2026, 9, 23, 20, 30, tzinfo=timezone.utc)}
    check("'эртага' is the day after it was sent, in Tashkent", due_day(request), date(2026, 9, 25))

    orders = [
        {"DocEntry": 1, "DocDate": "2026-09-25", "DocTotal": 100.5, "CANCELED": "N"},
        {"DocEntry": 2, "DocDate": "2026-09-25T00:00:00", "DocTotal": "200"},
        {"DocEntry": 3, "DocDate": "2026-09-25", "DocTotal": 999, "CANCELED": "Y"},
        {"DocEntry": 4, "DocDate": "2026-09-24", "DocTotal": 50},
    ]
    fig = figures.documents_on(orders, date(2026, 9, 25), "USD", tool="orders", as_of=date(2026, 9, 26))
    check("yesterday's orders, cancelled left out", (fig.totals, fig.count, fig.capped), ({"USD": 30050}, 2, False))
    capped = figures.documents_on(orders * 25, date(2026, 9, 25), "USD", tool="orders", as_of=None)
    check_true("100 rows = the push limit = a lower bound", capped.capped)
    check("rows without the columns -> unknown, not zero",
          figures.documents_on([{"x": 1}], date(2026, 9, 25), "USD", tool="orders", as_of=None).status, "unknown_format")
    check("no rows -> no data", figures.documents_on([], date(2026, 9, 25), "USD", tool="orders", as_of=None).status, "no_data")

    stock = [{"ItemCode": "A", "WhsCode": "01", "StockValue": 1000}, {"ItemCode": "B", "WhsCode": "01", "OnHand": 2, "AvgPrice": 25.5}]
    check("stock value: SAP's own, else on hand x price", figures.inventory_value(stock, "USD", as_of=None).totals, {"USD": 105100})
    check("change only where both days have it", figures.change({"USD": 500, "UZS": 7}, {"USD": 200}), {"USD": 300})


def test_receivables_rendering() -> None:
    """The receivables alert renders in both the empty and populated cases."""
    print("receivables rendering")
    from integrations.common.agent_loader import load_agent
    from integrations.sap.models import ARAging, ARInvoice

    receivables = load_agent("receivables")

    clean = ARAging(snapshot_date=date(2026, 8, 18), total_open_tiyin=10_000_000)
    text = receivables.render(clean, min_days=1)
    check_true("clean state is green", text.startswith("🟢"))

    aging = ARAging(snapshot_date=date(2026, 8, 18))
    aging.invoices = [
        ARInvoice(
            doc_entry=i,
            doc_num=1000 + i,
            card_code=f"C{i:03d}",
            card_name=f"Customer {i}",
            days_overdue=days,
            aging_bucket=aging_bucket(days),
            doc_total_tiyin=100_000_000,
            balance_due_tiyin=100_000_000,
            sales_person_name="Karimov B." if i % 2 else None,
        )
        for i, days in enumerate([5, 45, 75, 120, 200], start=1)
    ]
    aging.total_open_tiyin = 500_000_000
    aging.total_overdue_tiyin = 500_000_000
    aging.bucket_totals_tiyin = {"90_plus": 200_000_000}

    text = receivables.render(aging, min_days=1)
    check_true("headline critical", text.startswith("🔴"))
    check_true("age ranges folded into the headline", "90+ кун — 2" in text and "1–30 кун — 1" in text)
    check_true("two lines only", len(text.splitlines()) == 2)
    check_true("no per-invoice detail", "Customer" not in text)
    check_true("no owner breakdown", "Mas'ul xodim" not in text)
    one_range = receivables.render(aging, min_days=100)
    check_true("single range reads 'ичида'", "2 та ҳисоб-фактура, 90+ кун ичида" in one_range)
    check_true("the alert is Uzbek Cyrillic", latin_words(text) == [])

    text_30 = receivables.render(aging, min_days=30)
    check_true("min_days filters the 5-day invoice", "1–30 кун" not in text_30)


def test_daily_report_kpi() -> None:
    """Parsing employees' reported numbers, and how they render against target."""
    print("daily report KPI")
    from integrations.org_bot import kpi

    # 2026-09-16: nobody is asked for numbers — every role, sales included,
    # just reports what they did today.
    check("B2B Sotuv is not asked for numbers", kpi.metrics_for_role("b2b_sotuv"), ())
    check("Garmin Sotuv is not asked for numbers", kpi.metrics_for_role("garmin_sotuv"), ())
    check_true(
        "B2B Sotuv's ask has no numbers block",
        "Қўнғироқлар" not in kpi.build_request_text("Dmitriy", kpi.metrics_for_role("b2b_sotuv")),
    )
    check_true(
        "reminder has no numbers block",
        "Қўнғироқлар" not in kpi.build_reminder_text(kpi.metrics_for_role("garmin_sotuv")),
    )

    # The parser and formatter stay tested against the (currently unassigned)
    # sales set, so switching numbers back on for a role is safe.
    sales = kpi.SALES_METRICS
    check("sales metric set still defined", len(sales), 4)

    check(
        "labelled Uzbek reply",
        kpi.parse_metrics("Bugun uchrashuvlar 3, qo'ng'iroqlar 22, KP 5, yangi lidlar 2", sales),
        {"meetings": 3, "calls": 22, "proposals": 5, "new_leads": 2},
    )
    check(
        "number before the label, and a different apostrophe",
        kpi.parse_metrics("3 ta uchrashuv, 22 ta qoʻngʻiroq", sales),
        {"meetings": 3, "calls": 22},
    )
    check(
        "Russian labels",
        kpi.parse_metrics("встречи: 4, звонки: 30", sales),
        {"meetings": 4, "calls": 30},
    )
    check(
        "bare positional line",
        kpi.parse_metrics("3/20/5/2", sales),
        {"meetings": 3, "calls": 20, "proposals": 5, "new_leads": 2},
    )
    # A report with no numbers must never be given invented ones — the
    # Director's scorecard has to distinguish "nothing reported" from "zero".
    check("prose with no numbers stays empty", kpi.parse_metrics("Bugun ofisda ishladim.", sales), {})
    check("wrong count of bare numbers is not positional", kpi.parse_metrics("3 20", sales), {})
    check("no metrics for this role means no parsing", kpi.parse_metrics("uchrashuv 3", ()), {})

    rendered = kpi.format_metrics({"meetings": 5, "calls": 10}, sales)
    check_true("hit target marked", "Учрашувлар: 5/4 ✅" in rendered)
    check_true("under target marked", "Қўнғироқлар: 10/15 ⚠️" in rendered)
    check_true("unreported metric shows a dash", "Юборилган КП: —" in rendered)
    check("no metrics renders nothing", kpi.format_metrics({}, ()), "")
    check(
        "missing metrics are named",
        kpi.missing_metrics({"meetings": 5}, sales),
        ["Қўнғироқлар", "Юборилган КП", "Янги лидлар"],
    )

    ask = kpi.build_request_text("Dmitriy", sales)
    check_true("ask names the person", "Dmitriy" in ask)
    check_true("ask shows the numbers format", "Қўнғироқлар" in ask)
    check_true("text-only role gets no numbers block", "Қўнғироқлар" not in kpi.build_request_text("Aziz", ()))
    check(
        "Cyrillic labels are parsed too",
        kpi.parse_metrics("3 та учрашув, 22 та қўнғироқ", sales),
        {"meetings": 3, "calls": 22},
    )
    check_true("the ask is Uzbek Cyrillic", latin_words(kpi.build_request_text("Алишер", ())) == [])
    check_true("the reminder is Uzbek Cyrillic", latin_words(kpi.build_reminder_text(())) == [])


def test_permissions() -> None:
    """Written permission requests: intake wording, amounts, question order."""
    print("permissions (EMJ-SOP-ADM-01)")
    from integrations.org_bot import permissions

    # Starting a request
    check_true("command starts a request", permissions.wants_permission("/ruxsat"))
    check_true("Cyrillic request starts one", permissions.wants_permission("Рухсат керак: янги принтер"))
    check_true("Latin request starts one", permissions.wants_permission("ruxsat kerak, mehmonxonaga borish"))
    check_true("Russian request starts one", permissions.wants_permission("прошу разрешение на закупку"))
    # A mention is not a request: this one must reach the Director as a normal
    # message, not silently open a form.
    check_true("mentioning permission is not a request", not permissions.wants_permission("Директор рухсат берди"))
    check_true("unrelated message is not a request", not permissions.wants_permission("Бугун омборда ишладим"))

    # Amounts
    check("plain sum", permissions.parse_amount("5 000 000 сўм"), (500000000, "UZS"))
    check("dotted sum", permissions.parse_amount("5.000.000 so'm"), (500000000, "UZS"))
    check("dollars", permissions.parse_amount("$1200"), (120000, "USD"))
    check("no cost", permissions.parse_amount("0"), (0, "UZS"))
    check("unparseable stays unparsed", permissions.parse_amount("ҳали аниқ эмас"), (None, "UZS"))
    check("zero renders as zero", permissions.format_amount(0, "UZS"), "0")
    check("sum is grouped", permissions.format_amount(500000000, "UZS"), "5 000 000 сўм")
    check("raw answer is kept when unparsed", permissions.format_amount(None, "UZS", "тахминан 2 млн"), "тахминан 2 млн")

    # Question order follows the SOP form, and an unparseable amount still counts
    draft: dict = {}
    check("first question is the full name", permissions.next_missing_field(draft).key, "requester_full_name")
    draft["requester_full_name"] = "Алишер Каримов"
    check("then the position", permissions.next_missing_field(draft).key, "requester_position")
    draft["requester_position"] = "IT мутахассис"
    check("then the department", permissions.next_missing_field(draft).key, "department")
    draft["department"] = "IT"
    check("then the subject", permissions.next_missing_field(draft).key, "subject")
    draft["subject"] = "Принтер сотиб олиш"
    check("then the reason", permissions.next_missing_field(draft).key, "reason")
    draft["reason"] = "Эскиси ишламайди"
    check("then the amount", permissions.next_missing_field(draft).key, "amount")
    draft["amount_raw"] = "тахминан 2 млн"
    check("unparsed amount is not re-asked", permissions.next_missing_field(draft).key, "execute_by")
    for key, value in (
        ("execute_by", "25.09.2026"),
        ("decision_needed_by", "23.09.2026 12:00"),
        ("urgency", "оддий"),
        ("attachments", "йўқ"),
    ):
        draft[key] = value
    check("complete form asks nothing", permissions.next_missing_field(draft), None)

    # Buttons — with a REAL 36-character id. A short fake id is exactly what
    # hid the 65-byte "approved_conditional" button that made Telegram reject
    # every approver card (2026-09-21).
    real_id = "3fa85f64-5717-4562-b3fc-2c963f66afa6"
    buttons = [b for row in permissions.decision_keyboard(real_id)["inline_keyboard"] for b in row]
    check(
        "four SOP outcomes",
        [permissions.DECISION_CODES[b["callback_data"].split(":")[1]] for b in buttons],
        ["approved", "approved_conditional", "rejected", "info_needed"],
    )
    confirm = [b for row in permissions.confirm_keyboard(real_id)["inline_keyboard"] for b in row]
    for button in buttons + confirm:
        check_true(
            f"{button['callback_data'].rsplit(':', 1)[0]} button fits Telegram's 64-byte limit with a real id",
            len(button["callback_data"].encode()) <= 64,
        )

    card = permissions.request_card({**draft, "request_no": "EMJ-2026-0007", "amount_tiyin": None}, for_approver=True)
    check_true("card carries the request number", "EMJ-2026-0007" in card)
    check_true("card uses the SOP's own field wording", "Нимага рухсат сўралади" in card)

    # 2026-09-21: an unescaped "&" or "<" in a typed answer made Telegram
    # reject the approver's card, so the Director silently received nothing.
    risky = {**draft, "subject": "A&B <тест>", "requester_full_name": "Ali <IT>", "request_no": "EMJ-2026-0008"}
    risky_card = permissions.request_card(risky, for_approver=True)
    check_true("typed '&' is escaped in the card", "A&amp;B" in risky_card)
    check_true("typed '<' is escaped in the card", "&lt;тест&gt;" in risky_card and "<тест>" not in risky_card)
    check_true("requester name is escaped", "Ali &lt;IT&gt;" in risky_card)
    decided = permissions.decision_text({"status": "rejected", "request_no": "X", "decision_note": "нарх > лимит"})
    check_true("decision note is escaped", "нарх &gt; лимит" in decided)
    check_true("/bekor cancels", permissions.is_cancel(" /bekor ") and permissions.is_cancel("/бекор"))
    check_true("an answer is not a cancel", not permissions.is_cancel("бекор қилинган буюртма"))


def test_task_tracker() -> None:
    """A3 deadlines/reminders and the plan's И coefficients (A1, A3)."""
    print("task tracker (A3) and И")
    from datetime import datetime, timezone

    from integrations.org_bot import task_tracker as tt
    from integrations.org_bot.prompt import build_classify_message

    today = date(2026, 9, 23)  # a Wednesday
    check("stated deadline kept", tt.parse_due_date("2026-09-25", today), date(2026, 9, 25))
    check("today is a valid deadline", tt.parse_due_date("2026-09-23", today), today)
    check("no deadline stays none", tt.parse_due_date(None, today), None)
    check("a past date is dropped, not guessed", tt.parse_due_date("2026-09-01", today), None)
    check("garbage is dropped", tt.parse_due_date("juma", today), None)
    check("a year-plus date is a misread", tt.parse_due_date("2028-01-01", today), None)

    check("'Ertaga' button", tt.due_from_choice("1", today), (True, date(2026, 9, 24)))
    check("'Muddatsiz' button is no deadline", tt.due_from_choice("n", today), (True, None))
    check("unknown button refused", tt.due_from_choice("x", today)[0], False)
    buttons = [b for row in tt.deadline_keyboard(987654321)["inline_keyboard"] for b in row]
    check_true("deadline buttons fit Telegram's 64 bytes", all(len(b["callback_data"].encode()) <= 64 for b in buttons))
    check_true("card line names the weekday", "жума" in tt.deadline_line(date(2026, 9, 25), today))
    check_true("card line says 'эртага'", "эртага" in tt.deadline_line(date(2026, 9, 24), today))

    message = build_classify_message("ertaga hisobotni tayyorla", today=today)
    check_true("classifier is told today's date", "Wednesday, 2026-09-23" in message)

    def at(day: int, hour: int) -> datetime:
        return datetime(2026, 9, day, hour - 5, 0, tzinfo=timezone.utc)  # Tashkent = UTC+5

    tasks = [
        {"display_name": "Aziz", "task_summary": "A", "status": "done", "due_date": date(2026, 9, 22), "completed_at": at(22, 17)},
        {"display_name": "Aziz", "task_summary": "B", "status": "done", "due_date": date(2026, 9, 21), "completed_at": at(22, 10)},
        {"display_name": "Dilnoza", "task_summary": "C", "status": "sent", "due_date": date(2026, 9, 22), "completed_at": None},
        {"display_name": "Dilnoza", "task_summary": "D", "status": "sent", "due_date": today, "completed_at": None},
        {"display_name": "Dilnoza", "task_summary": "E", "status": "sent", "due_date": None, "completed_at": None},
    ]
    score = tt.score_tasks(tasks, date(2026, 9, 21), today)
    check("tasks judged (open, due today, not yet late)", score.due, 3)
    check("done by 17:00 on the day counts as on time", score.on_time, 1)
    check("done the day after is late", score.late, 1)
    check("open past deadline is overdue", score.open_overdue, 1)
    check("A3 И", round(score.index, 2), 0.33)

    rows = [
        {"display_name": "Aziz", "status": "submitted", "report_date": today, "submitted_at": at(23, 16)},
        {"display_name": "Bobur", "status": "submitted", "report_date": today, "submitted_at": at(23, 19)},
        {"display_name": "Dilnoza", "status": "asked", "report_date": today, "submitted_at": None},
    ]
    reports = tt.score_reports(rows)
    check("reports: 2 of 3 sent", (reports.reported, reports.asked), (2, 3))
    check("a report after 18:00 isn't on time", reports.on_time, 1)
    check("A1 И = (2/3)*(1/2)", round(reports.index, 2), 0.33)

    text = tt.weekly_text(date(2026, 9, 21), today, score, reports, {"approved": 2, "submitted": 1})
    check_true("scorecard shows the tasks И", "1/3 ўз вақтида" in text and "И = 0.33" in text)
    check_true("below 0.50 is red", "🔴" in text)
    check_true("names who didn't report", "Dilnoza (1 кун)" in text)
    check_true("scorecard is Uzbek Cyrillic", latin_words(text, allow={"Aziz", "Bobur", "Dilnoza", "C"}) == [])
    check_true(
        "reminder is Uzbek Cyrillic",
        latin_words(tt.reminder_text({"due_date": today, "task_summary": "Ҳисобот"}, today)) == [],
    )
    check_true("names the overdue task", "Dilnoza — C" in text)
    check_true("permissions counted", "тасдиқланди 2" in text and "кутилмоқда 1" in text)
    quiet = tt.weekly_text(date(2026, 9, 21), today, tt.TaskScore(), None, None)
    check_true("no reports section when reports are off", "hisobot" not in quiet)
    check_true("typed names are escaped", "&lt;" in tt.overdue_director_text(
        [{"display_name": "A<b>", "task_summary": "x", "due_date": today}]
    ))


def test_names_and_routing() -> None:
    """Every employee's typed name, and the Director addressing one person."""
    print("names and person routing")
    from integrations.org_bot import names
    from integrations.org_bot.ops_manager import _pick_person, _plain
    from integrations.org_bot.prompt import build_classify_message

    worker = {"role": "it", "full_name": None, "display_name": "GMHRD"}
    check_true("an employee without a name must give one", names.needs_name(worker))
    check_true("the Director is never stopped for a name", not names.needs_name({"role": "operatsion_direktor"}))
    check_true("a saved name is enough", not names.needs_name({**worker, "full_name": "Алишер Каримов"}))
    check("Telegram name only as a fallback", names.person_name(worker), "GMHRD")
    check("typed name wins", names.person_name({**worker, "full_name": "Алишер Каримов"}), "Алишер Каримов")
    check_true("the name question is Uzbek Cyrillic", latin_words(names.ASK_TEXT) == [])

    roster = {"E1": {"role": "it", "full_name": "Алишер Каримов"}, "E2": {"role": "ombor", "full_name": "Бобур Алиев"}}
    result = {"target_type": "employee", "target_role": "b2b_sotuv", "target_employee": "e2"}
    person, known = _pick_person(result, roster)
    check_true("named person found", known and person is roster["E2"])
    check("their own department wins over the model's guess", result["target_role"], "ombor")
    check("no person named -> whole department", _pick_person({"target_employee": None}, roster), (None, True))
    check("unknown person -> ask, never broadcast", _pick_person({"target_employee": "E9"}, roster), (None, False))

    message = build_classify_message("Alisherga ayt", roster=["E1 = Алишер Каримов (it)"])
    check_true("classifier sees the employee list", "E1 = Алишер Каримов (it)" in message)
    check("preview drops the AI's tags", _plain("<b>Отчёт</b>  тайёрланг"), "Отчёт тайёрланг")


def test_report_or_message() -> None:
    """A report never needs Telegram's reply; an unclear message gets one tap."""
    print("report or message")
    from integrations.org_bot.ops_manager import report_message_kind, report_or_relay_keyboard

    check("plain message, no task -> report", report_message_kind(False, False, 0), "report")
    check("reply to the ask/reminder -> report, even with tasks", report_message_kind(True, False, 3), "report")
    check("reply to a task card -> task update", report_message_kind(False, True, 1), "task_update")
    check("plain message with a task in flight -> ask", report_message_kind(False, False, 1), "ask")
    check("plain message with several tasks -> ask", report_message_kind(False, False, 4), "ask")
    real_id = "3fa85f64-5717-4562-b3fc-2c963f66afa6"
    buttons = [b for row in report_or_relay_keyboard(real_id)["inline_keyboard"] for b in row]
    check("three choices", [b["callback_data"].split(":")[0] for b in buttons], ["asrep", "relayok", "relayno"])
    check_true("every button fits Telegram's 64 bytes", all(len(b["callback_data"].encode()) <= 64 for b in buttons))
    check_true("buttons are Uzbek Cyrillic", all(latin_words(b["text"]) == [] for b in buttons))


def test_db_viewer() -> None:
    """The /db viewer: locked by default, read-only, escapes what it shows."""
    print("database viewer")
    import base64
    from datetime import datetime, timezone

    from fastapi.testclient import TestClient
    from pydantic import SecretStr

    from integrations.api import app as api_app
    from integrations.api import db_viewer
    from integrations.common.config import settings

    fake_rows = {
        "relations": [{"name": "employees", "kind": "r"}, {"name": "v_ar_aging_latest", "kind": "v"}],
        "columns": [{"column_name": "id"}, {"column_name": "full_name"}, {"column_name": "created_at"}],
        "rows": [{"id": 7, "full_name": "<script>x</script>", "created_at": datetime(2026, 9, 26, 3, 0, tzinfo=timezone.utc)}],
    }
    seen_queries: list[str] = []

    async def fake_read(query, params=None):
        text = query if isinstance(query, str) else repr(query)
        seen_queries.append(text)
        if "pg_class" in text:
            return fake_rows["relations"]
        if "information_schema.columns" in text:
            return fake_rows["columns"]
        if "count(*)" in text:
            return [{"n": 1}]
        return fake_rows["rows"]

    original_read, original_password = db_viewer._read, settings.db_viewer_password
    db_viewer._read = fake_read
    client = TestClient(api_app.app)
    try:
        settings.db_viewer_password = SecretStr("")
        check("switched off without a password", client.get("/db").status_code, 404)

        settings.db_viewer_password = SecretStr("s3cret-long")
        check("asks for a login", client.get("/db").status_code, 401)
        wrong = {"Authorization": "Basic " + base64.b64encode(b"admin:nope").decode()}
        check("wrong password refused", client.get("/db", headers=wrong).status_code, 401)

        good = {"Authorization": "Basic " + base64.b64encode(b"any:s3cret-long").decode()}
        index = client.get("/db", headers=good)
        check("right password lets you in", index.status_code, 200)
        check_true("lists the tables", "employees" in index.text and "1 қатор" in index.text)
        check_true("never indexed or cached", index.headers.get("x-robots-tag", "").startswith("noindex")
                   and index.headers.get("cache-control") == "no-store")

        page = client.get("/db/employees", headers=good)
        check_true("shows a row", page.status_code == 200 and "full_name" in page.text)
        check_true("what the data holds is escaped", "&lt;script&gt;" in page.text and "<script>x" not in page.text)
        check_true("timestamps in Tashkent time", "2026-09-26 08:00:00" in page.text)
        check_true("newest first", any("ORDER BY" in q and "created_at" in q for q in seen_queries))

        seen_queries.clear()
        missing = client.get("/db/pg_authid", headers=good)
        check_true("a table outside the list is refused", "Бундай жадвал йўқ" in missing.text)
        check_true("...without ever being queried", not any("pg_authid" in q for q in seen_queries))
    finally:
        db_viewer._read = original_read
        settings.db_viewer_password = original_password


def test_plan_agents() -> None:
    """B2 cash calendar, B4 data quality, E1 monthly KPI."""
    print("B2 / B4 / E1")
    from datetime import datetime, timezone

    from integrations.common.agent_loader import load_agent
    from integrations.org_bot import task_tracker as tt

    # ---- B2: 30-day cash calendar
    cc = load_agent("cash-calendar")
    today = date(2026, 9, 28)
    sent = datetime(2026, 9, 27, 5, tzinfo=timezone.utc)
    invoices = [
        {"due_date": date(2026, 9, 30), "balance_due_tiyin": 320000, "currency": "USD"},
        {"due_date": date(2026, 10, 20), "balance_due_tiyin": 550000, "currency": "USD"},
        {"due_date": date(2026, 9, 10), "balance_due_tiyin": 738436, "currency": "USD"},
        {"due_date": None, "balance_due_tiyin": 5, "currency": "USD"},
    ]
    payments = [
        {"request_no": "EMJ-2026-0007", "subject": "Принтер", "amount_tiyin": 1_500_000_000, "currency": "UZS",
         "execute_by": "30.09.2026", "submitted_at": sent},
        {"request_no": "EMJ-2026-0008", "subject": "Мебель", "amount_tiyin": 100, "currency": "UZS",
         "execute_by": "тезроқ", "submitted_at": sent},
        {"request_no": "EMJ-2026-0009", "subject": "Кеча", "amount_tiyin": 100, "currency": "UZS",
         "execute_by": "01.09.2026", "submitted_at": sent},
    ]
    cal = cc.build(invoices, payments, today, capped=False)
    check("four blocks cover the 30 days", (cal.incoming[0].start, cal.incoming[-1].end), (today, date(2026, 10, 27)))
    check("invoice due this week lands in week 1", (cal.incoming[0].count, cal.incoming[0].totals), (1, {"USD": 320000}))
    check("overdue kept apart", (cal.overdue_count, cal.overdue), (1, {"USD": 738436}))
    check("only in-window approved payments go out", [p["no"] for p in cal.outgoing], ["EMJ-2026-0007"])
    check("an unreadable date is counted, not placed", cal.unclear, 1)
    text = cc.render(cal)
    check_true("says there is no balance", "Касса қолдиғи уланмаган" in text)
    check_true("the calendar is Uzbek Cyrillic", latin_words(text, allow={"EMJ"}) == [])
    check_true("a capped invoice feed is a lower bound", "камида" in cc.render(cc.build(invoices, [], today, capped=True)))

    # ---- B4: data quality
    dq = load_agent("data-quality")
    inp = dq.Inputs(today=today)
    inp.invoices = [
        {"doc_num": 2253, "doc_date": date(2026, 9, 1), "due_date": date(2026, 9, 10), "currency": "USD",
         "doc_total_tiyin": 100, "balance_due_tiyin": 100, "sales_person_code": -1},
        {"doc_num": 2332, "doc_date": date(2026, 9, 5), "due_date": None, "currency": "USD",
         "doc_total_tiyin": 100, "balance_due_tiyin": 150, "sales_person_code": 4},
    ]
    inp.feeds = {"invoices": {"at": datetime(2026, 9, 28, 2, tzinfo=timezone.utc), "rows": 20},
                 "payments": {"at": datetime(2026, 9, 28, 2, tzinfo=timezone.utc), "rows": 100}}
    inp.unnamed = ["GMHRD"]
    report = dq.render(inp)
    check_true("no sales person is flagged", "Масъул сотувчиси йўқ очиқ ҳисоб-фактура: 1 та" in report and "#2253" in report)
    check_true("missing due date is flagged", "Тўлов муддати йўқ ҳисоб-фактура: 1 та (#2332)" in report)
    check_true("a balance above the total is flagged", "Қолдиғи нотўғри ҳисоб-фактура: 1 та (#2332)" in report)
    check_true("a capped feed is flagged", "тўловлар — чекланган (100 та" in report)
    check_true("a feed that never came is flagged", "буюртмалар — ҳеч қачон келмаган" in report)
    check_true("unnamed employees are flagged", "Исм ёзмаган ходимлар: 1" in report)
    check_true("the report is Uzbek Cyrillic", latin_words(report, allow={"GMHRD"}) == [])
    clean = dq.Inputs(today=today, feeds={t: {"at": datetime(2026, 9, 28, 2, tzinfo=timezone.utc), "rows": 1}
                                          for t in dq.FEED_LABELS})
    check_true("clean data says so", dq.render(clean).count("✅ тоза") == 2)

    # ---- E1: monthly KPI
    def at(d: int, h: int) -> datetime:
        return datetime(2026, 9, d, h - 5, tzinfo=timezone.utc)

    reports = [{"display_name": "Алишер", "status": "submitted", "report_date": date(2026, 9, d), "submitted_at": at(d, 17)}
               for d in range(1, 11)]
    reports += [{"display_name": "Дилноза", "status": "asked" if d % 2 else "submitted", "report_date": date(2026, 9, d),
                 "submitted_at": at(d, 17)} for d in range(1, 11)]
    tasks = [{"display_name": "Бобур", "task_summary": "x", "status": "done", "due_date": date(2026, 9, 10),
              "completed_at": at(12, 10)}]
    kpis = tt.employee_kpis(reports, tasks, date(2026, 9, 1), date(2026, 9, 30))
    check("worst first: late tasks, then half the reports, then all good",
          [(k.name, k.mark) for k in kpis], [("Бобур", "🔴"), ("Дилноза", "🟡"), ("Алишер", "🟢")])
    monthly = tt.monthly_text(date(2026, 9, 1), kpis)
    check_true("month named in Uzbek", "сентябр 2026" in monthly)
    check_true("per-person line", "Дилноза — ҳисобот 5/10 (5) · топшириқ —" in monthly)
    check_true("the KPI message is Uzbek Cyrillic", latin_words(monthly) == [])
    check_true("an empty month says so", "на ҳисобот сўралди" in tt.monthly_text(date(2026, 9, 1), []))


def test_payment_gate() -> None:
    """B1: requests route by amount to whoever holds that limit."""
    print("payment gate (B1)")
    from integrations.org_bot import permissions as p

    tiers = p.parse_tiers("20 000 000:222, 5000000:111, junk, 0:9")
    check("limits parsed, sorted, junk skipped", tiers, [(500000000, 111), (2000000000, 222)])
    check("empty setting = no tiers", p.parse_tiers(""), [])
    check("small amount -> first limit", p.tier_approver(300000000, "UZS", tiers), 111)
    check("exactly the limit stays in that tier", p.tier_approver(500000000, "UZS", tiers), 111)
    check("mid amount -> second limit", p.tier_approver(1500000000, "UZS", tiers), 222)
    check("above every limit -> Director", p.tier_approver(9000000000, "UZS", tiers), None)
    check("dollars always go to the Director", p.tier_approver(100000, "USD", tiers), None)
    check("no cost goes to the Director", p.tier_approver(0, "UZS", tiers), None)
    check("unreadable amount goes to the Director", p.tier_approver(None, "UZS", tiers), None)
    check("no tiers configured -> Director", p.tier_approver(300000000, "UZS", []), None)


def test_permission_form() -> None:
    """The company's own SOP .docx is filled in place — never regenerated."""
    print("permission form (company docx filled in place)")
    from datetime import datetime, timezone

    from docx import Document

    from integrations.org_bot import docx_form

    check("blank after label is filled", docx_form.fill_blank_after("Бўлим: ______", "Бўлим", "IT"), "Бўлим: IT")
    check("empty answer leaves the blank", docx_form.fill_blank_after("Бўлим: ______", "Бўлим", ""), "Бўлим: ______")
    check(
        "only the chosen box is ticked",
        docx_form.tick_decision("☐ Тасдиқланди     ☐ Шарт билан тасдиқланди", "approved_conditional"),
        "☐ Тасдиқланди     ☑ Шарт билан тасдиқланди",
    )

    request = {
        "request_no": "EMJ-2026-0009",
        "requester_name": "tg_profile_name",
        "requester_full_name": "Алишер Каримов",
        "requester_position": "IT мутахассис",
        "department": "IT",
        "submitted_to": "Операцион директор",
        "subject": "Принтер сотиб олиш",
        "reason": "Эскиси ишламайди, ҳужжатлар кечикяпти; офис учун янги лазерли принтер олиш ва эскисини "
        "омборга топшириш таклиф қилинади",
        "amount_raw": "2 000 000 сўм",
        "execute_by": "25.09.2026",
        "decision_needed_by": "23.09.2026 12:00",
        "urgency": "оддий",
        "attachments": "йўқ",
        "requester_telegram_user_id": 111,
        "submitted_at": datetime(2026, 9, 21, 6, 0, tzinfo=timezone.utc),
        "status": "approved_conditional",
        "approved_terms": "1 800 000 сўм, 30.09.2026 гача",
        "decision_note": "Фақат шартнома билан",
        "decided_by": "Бобур Алиев, Операцион директор",
        "decided_by_telegram_user_id": 222,
        "decided_at": datetime(2026, 9, 21, 8, 30, tzinfo=timezone.utc),
    }
    source = [p.text for p in Document(str(docx_form.TEMPLATE_PATH)).paragraphs]
    filled = [p.text for p in Document(str(docx_form.fill_form(request))).paragraphs]
    body = chr(10).join(filled)
    check("same paragraphs as the company's file", len(filled), len(source))
    title_at = next(i for i, t in enumerate(source) if docx_form.FORM_TITLE in t)
    check("page 1 is untouched (incl. director's approval line)", filled[: title_at + 1], source[: title_at + 1])
    for expected in (
        "EMJ-2026-0009", "21.09.2026 11:00", "Алишер Каримов, IT мутахассис", "Принтер сотиб олиш",
        "2 000 000 сўм", "Фақат шартнома билан", "Бобур Алиев, Операцион директор", "21.09.2026 13:30",
    ):
        check_true(f"form contains '{expected}'", expected in body)
    check_true("no Telegram profile name on the form", "tg_profile_name" not in body)
    form_part = chr(10).join(filled[title_at:])
    check_true("no electronic signature text", "Telegram ID" not in form_part and "электрон" not in form_part)
    check_true("employee signature left for hand", any(t.startswith("Иловалар") and t.rstrip().endswith("_") for t in filled))
    check_true("approver signature left for hand", any(t.startswith("Имзо ___") for t in filled))
    reason_at = next(i for i, t in enumerate(filled) if t.startswith("Сабаб ва таклиф"))
    check_true("long reason continues on the form's next line", "омборга топшириш" in filled[reason_at + 1])
    check_true("reason line keeps its width", len(filled[reason_at]) <= len(source[reason_at]))
    check_true("the conditional box is ticked", "☑ Шарт билан тасдиқланди" in body)
    check_true("the other boxes stay empty", "☐ Тасдиқланди" in body and "☐ Рад этилди" in body)
    changed = [i for i, (a, b) in enumerate(zip(source, filled)) if a != b]
    check_true("only form lines changed", all(i > title_at for i in changed))


def main() -> int:
    """Run every check.

    Returns:
        0 if all checks pass, 1 otherwise.
    """
    for suite in (
        test_references,
        test_bot_flows,
        test_money,
        test_time,
        test_aging,
        test_telegram,
        test_org_bot,
        test_permissions,
        test_permission_form,
        test_task_tracker,
        test_payment_gate,
        test_plan_agents,
        test_db_viewer,
        test_names_and_routing,
        test_report_or_message,
        test_daily_report_kpi,
        test_brief_rendering,
        test_dates_and_figures,
        test_receivables_rendering,
    ):
        suite()

    print()
    if FAILURES:
        print(f"{len(FAILURES)} check(s) FAILED: {', '.join(FAILURES)}")
        return 1
    print("All checks passed.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
